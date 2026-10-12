import zipfile

from db import ActividadController, PDFService
from constants import TIPOS_ACTIVIDAD_PREDETERMINADOS
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pantallas.actividades.controller import generar_plantilla_actividad


class _Auth:
    def __init__(self, rol="propietario"):
        self._rol = rol

    def puede_escribir(self):
        return self._rol in ("propietario", "administrador")

    def obtener_usuario_actual(self):
        return {"rol": self._rol}

    @staticmethod
    def mensaje_sin_permiso():
        return "Sin permiso."


class _TipoDAO:
    def __init__(self, tipos=()):
        self.tipos = [dict(tipo) for tipo in tipos]
        self.siguiente_id = max((tipo["id_tipo"] for tipo in self.tipos), default=0) + 1

    def obtener_todos(self):
        return sorted(self.tipos, key=lambda tipo: tipo["nombre"])

    def crear(self, codigo=None, nombre=None):
        tipo = {"id_tipo": self.siguiente_id, "nombre": nombre or codigo}
        self.siguiente_id += 1
        self.tipos.append(tipo)

    def actualizar(self, id_tipo, codigo=None, nombre=None):
        tipo = next(tipo for tipo in self.tipos if tipo["id_tipo"] == id_tipo)
        tipo["nombre"] = nombre or codigo


def _controller(tipos=(), rol="propietario"):
    controller = ActividadController.__new__(ActividadController)
    controller.auth = _Auth(rol)
    controller.tipo_dao = _TipoDAO(tipos)
    return controller


def test_listar_tipos_conserva_el_catalogo_completo():
    personalizado = {"id_tipo": 20, "nombre": "Jornada comunitaria"}
    existente = {"id_tipo": 21, "nombre": "Sesiones educativas"}
    controller = _controller([personalizado, existente])

    tipos = controller.listar_tipos_actividad()

    assert tipos == [personalizado, existente]


def test_listar_no_recrea_nombre_predeterminado_tras_renombrarlo():
    renombrado = {"id_tipo": 1, "nombre": "Sesiones de formación"}
    controller = _controller([renombrado])

    tipos = controller.listar_tipos_actividad()

    assert tipos == [renombrado]


def test_listar_catalogo_vacio_inicializa_los_tipos_predeterminados():
    controller = _controller()

    tipos = controller.listar_tipos_actividad()

    assert {tipo["nombre"] for tipo in tipos} == {
        nombre for _, nombre in TIPOS_ACTIVIDAD_PREDETERMINADOS
    }


def test_gestion_tipo_actividad_agrega_edita_y_rechaza_nombres_duplicados():
    controller = _controller([{"id_tipo": 1, "nombre": "Tipo existente"}])

    ok, _ = controller.crear_tipo_actividad("  Nueva actividad  ")
    assert ok
    creado = next(tipo for tipo in controller.tipo_dao.tipos if tipo["id_tipo"] == 2)
    assert creado["nombre"] == "Nueva actividad"

    ok, mensaje = controller.crear_tipo_actividad(" nueva   actividad ")
    assert not ok
    assert "Ya existe" in mensaje

    ok, _ = controller.actualizar_tipo_actividad(2, "Nombre actualizado")
    assert ok
    assert creado["nombre"] == "Nombre actualizado"


def test_gestion_tipo_actividad_no_permite_eliminar_o_escribir_sin_permiso():
    controller = _controller(
        [{"id_tipo": 1, "nombre": "No borrar"}],
        rol="lector",
    )

    ok, _ = controller.crear_tipo_actividad("Otra actividad")
    assert not ok
    ok, _ = controller.actualizar_tipo_actividad(1, "Renombrado")
    assert not ok
    assert controller.tipo_dao.tipos == [{"id_tipo": 1, "nombre": "No borrar"}]
    assert not hasattr(controller.tipo_dao, "desactivar")


def test_solo_propietario_puede_administrar_tipos_de_actividad():
    controller = _controller(rol="administrador")

    ok, mensaje = controller.crear_tipo_actividad("Tipo no permitido")

    assert not ok
    assert "Solo el usuario propietario" in mensaje
    assert controller.tipo_dao.tipos == []


def test_reporte_word_incluye_tipos_personalizados_y_sus_totales(tmp_path):
    ruta = tmp_path / "reporte_tipos_personalizados.docx"
    PDFService().generar_word(
        {
            "resumen_por_tipo": [
                {
                    "nombre": "Jornada comunitaria",
                    "total_actividades": 3,
                    "total_participantes": 27,
                }
            ],
        },
        str(ruta),
    )

    with zipfile.ZipFile(ruta) as archivo:
        documento = archivo.read("word/document.xml").decode("utf-8")

    assert "Jornada comunitaria" in documento
    assert ">3<" in documento
    assert ">27<" in documento


def test_plantilla_importacion_usa_catalogo_actual_para_validacion_de_tipo(tmp_path):
    ruta_entrada = tmp_path / "plantilla_base.xlsx"
    ruta_salida = tmp_path / "plantilla_generada.xlsx"
    workbook = Workbook()
    hoja = workbook.active
    validacion_tipos_fija = DataValidation(
        type="list",
        formula1='"Tipo antiguo"',
    )
    validacion_tipos_fija.add("I4:I1048576")
    hoja.add_data_validation(validacion_tipos_fija)
    validacion_estado = DataValidation(type="list", formula1='"Estado"')
    validacion_estado.add("B4:B1048576")
    hoja.add_data_validation(validacion_estado)
    workbook.save(ruta_entrada)

    nombres = [f"Tipo de actividad personalizado {indice}" for indice in range(30)]
    generar_plantilla_actividad(ruta_entrada, ruta_salida, nombres)

    resultado = load_workbook(ruta_salida)
    hoja_generada = resultado.active
    validaciones = list(hoja_generada.data_validations.dataValidation)
    validacion_generada = next(
        validacion for validacion in validaciones if str(validacion.sqref) == "I4:I1048576"
    )
    columna_opciones = hoja_generada.max_column
    letra_columna_opciones = hoja_generada.cell(row=1, column=columna_opciones).column_letter
    assert validacion_generada.formula1 == (
        f"=${letra_columna_opciones}$1:${letra_columna_opciones}$30"
    )
    assert any(str(validacion.sqref) == "B4:B1048576" for validacion in validaciones)
    assert hoja_generada.column_dimensions[letra_columna_opciones].hidden
    assert [
        hoja_generada.cell(row=fila, column=columna_opciones).value
        for fila in range(1, len(nombres) + 1)
    ] == nombres


def test_plantilla_usa_opciones_directas_y_muestra_flecha_desplegable(tmp_path):
    ruta_entrada = tmp_path / "plantilla_base.xlsx"
    ruta_salida = tmp_path / "plantilla_generada.xlsx"
    Workbook().save(ruta_entrada)
    nombres = ["Tipo uno", "Tipo dos"]

    generar_plantilla_actividad(ruta_entrada, ruta_salida, nombres)

    hoja = load_workbook(ruta_salida).active
    validacion = next(
        validacion
        for validacion in hoja.data_validations.dataValidation
        if str(validacion.sqref) == "I4:I1048576"
    )
    assert validacion.type == "list"
    assert validacion.formula1 == '"Tipo uno,Tipo dos"'
    assert validacion.showDropDown is False
    assert validacion.showErrorMessage is True


def test_no_se_genera_plantilla_si_no_hay_tipos_de_actividad(tmp_path):
    ruta_entrada = tmp_path / "plantilla_base.xlsx"
    ruta_salida = tmp_path / "plantilla_generada.xlsx"
    Workbook().save(ruta_entrada)

    try:
        generar_plantilla_actividad(ruta_entrada, ruta_salida, [])
    except ValueError as exc:
        assert "No hay tipos de actividad" in str(exc)
    else:
        raise AssertionError("Debió rechazar la generación sin tipos de actividad.")


def test_generar_plantilla_no_sobrescribe_el_archivo_original(tmp_path):
    ruta = tmp_path / "plantilla.xlsx"
    Workbook().save(ruta)

    try:
        generar_plantilla_actividad(ruta, ruta, ["Tipo vigente"])
    except ValueError as exc:
        assert "distinto al original" in str(exc)
    else:
        raise AssertionError("Debió rechazar la ruta de salida igual a la plantilla original.")
