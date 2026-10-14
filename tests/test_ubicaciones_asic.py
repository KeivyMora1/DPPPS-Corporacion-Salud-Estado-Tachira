from datetime import date

from openpyxl import Workbook, load_workbook

from db import (
    ASIC_PREDETERMINADOS,
    ACTIVIDAD_ESTABLECIMIENTO_OPCIONES,
    ActividadController,
    ExcelService,
    ExcelServiceVoceria,
    MUNICIPIOS_TACHIRA,
    PARROQUIAS_POR_MUNICIPIO_TACHIRA,
    ParroquiaDAO,
    UBICACIONES_ASIC_PREDETERMINADAS,
    VOCERIA_CENTRO_DISPENSADOR_OPCIONES,
    VoceriaController,
    build_voceria_template,
)
from pantallas.actividades.controller import generar_plantilla_actividad


def _ubicaciones():
    return [
        {"asic_nombre": asic, "municipio": municipio, "parroquia": parroquia}
        for asic, municipio, parroquia in UBICACIONES_ASIC_PREDETERMINADAS
    ]


def test_catalogo_geografico_cubre_los_29_municipios_y_66_parroquias():
    assert len(MUNICIPIOS_TACHIRA) == 29
    assert sum(map(len, PARROQUIAS_POR_MUNICIPIO_TACHIRA.values())) == 66
    assert MUNICIPIOS_TACHIRA == list(PARROQUIAS_POR_MUNICIPIO_TACHIRA)
    assert PARROQUIAS_POR_MUNICIPIO_TACHIRA["San Cristóbal"] == [
        "La Concordia",
        "Pedro María Morantes",
        "San Juan Bautista",
        "San Sebastián",
        "Francisco Romero Lobo",
    ]
    assert PARROQUIAS_POR_MUNICIPIO_TACHIRA["Panamericano"] == [
        "San Pablo",
        "La Palmita",
    ]
    assert PARROQUIAS_POR_MUNICIPIO_TACHIRA["Sucre"] == [
        "Sucre",
        "Eleazar López Contreras",
        "San Pablo",
    ]
    assert all(
        not parroquia.startswith("Parroquia ")
        for parroquias in PARROQUIAS_POR_MUNICIPIO_TACHIRA.values()
        for parroquia in parroquias
    )
    assert {
        asic for asic, _, _ in UBICACIONES_ASIC_PREDETERMINADAS
    } == {nombre for _, nombre in ASIC_PREDETERMINADOS}
    pares_oficiales = {
        (municipio, parroquia)
        for municipio, parroquias in PARROQUIAS_POR_MUNICIPIO_TACHIRA.items()
        for parroquia in parroquias
    }
    for _, nombre_asic in ASIC_PREDETERMINADOS:
        assert {
            (municipio, parroquia)
            for asic, municipio, parroquia in UBICACIONES_ASIC_PREDETERMINADAS
            if asic == nombre_asic
        } == pares_oficiales


def test_catalogo_geografico_no_depende_del_asic():
    dao = ParroquiaDAO()
    esperado = {
        (municipio, parroquia)
        for municipio, parroquias in PARROQUIAS_POR_MUNICIPIO_TACHIRA.items()
        for parroquia in parroquias
    }

    assert {
        (ubicacion["municipio"], ubicacion["parroquia"])
        for ubicacion in dao.listar_por_asic(1)
    } == esperado
    assert dao.listar_por_asic(1) == dao.listar_por_asic(99)


def test_schema_contiene_tablas_municipio_y_parroquia_con_seeds():
    with open("schema.sql", encoding="utf-8") as archivo:
        schema = archivo.read()

    assert "CREATE TABLE IF NOT EXISTS municipio" in schema
    assert "CREATE TABLE IF NOT EXISTS parroquia" in schema
    # Las tablas relacionales deben tener sus seeds
    assert "INSERT IGNORE INTO municipio" in schema
    assert "INSERT IGNORE INTO parroquia" in schema
    # La tabla parroquia ya no depende de id_asic
    assert "REFERENCES municipio(id_municipio)" in schema


def test_plantilla_de_actividades_usa_catalogo_geografico_oficial(tmp_path):
    entrada = tmp_path / "plantilla_base.xlsx"
    salida = tmp_path / "plantilla_actividades.xlsx"
    Workbook().save(entrada)

    generar_plantilla_actividad(entrada, salida, ["Tipo"], _ubicaciones())

    workbook = load_workbook(salida)
    hoja = workbook.active
    validaciones = {
        str(validacion.sqref): validacion
        for validacion in hoja.data_validations.dataValidation
    }
    assert validaciones["E4:E1048576"].formula1 == "=CatalogoASIC"
    assert validaciones["C4:C1048576"].formula1 == "=CatalogoMunicipios"
    assert validaciones["D4:D1048576"].formula1 == (
        '=INDIRECT("Parroquia"&MATCH(C4,CatalogoMunicipios,0))'
    )
    assert hoja["I4"].value is None
    assert workbook.defined_names["CatalogoASIC"].attr_text == (
        "'Catalogos'!$A$2:$A$26"
    )
    assert workbook.defined_names["CatalogoMunicipios"].attr_text == (
        "'Catalogos'!$B$2:$B$30"
    )
    assert workbook.defined_names["Parroquia1"].attr_text == (
        "'Catalogos'!$C$2:$C$2"
    )
    assert workbook.defined_names["Parroquia29"].attr_text == (
        "'Catalogos'!$AE$2:$AE$5"
    )
    hoja_listas = workbook["Catalogos"]
    assert hoja_listas.sheet_state == "visible"
    assert [hoja_listas.cell(row=fila, column=1).value for fila in range(2, 27)] == sorted({
        ubicacion["asic_nombre"] for ubicacion in _ubicaciones()
    })
    assert [hoja_listas.cell(row=fila, column=2).value for fila in range(2, 31)] == (
        sorted(MUNICIPIOS_TACHIRA)
    )
    assert sum(
        hoja_listas.cell(row=fila, column=columna).value is not None
        for columna in range(3, 32)
        for fila in range(2, 7)
    ) == 66


def test_plantilla_de_vocerias_usa_catalogo_geografico_oficial(tmp_path):
    salida = tmp_path / "plantilla_vocerias.xlsx"

    build_voceria_template(salida, _ubicaciones())

    workbook = load_workbook(salida)
    hoja = workbook.active
    validaciones = {
        str(validacion.sqref): validacion
        for validacion in hoja.data_validations.dataValidation
    }
    assert validaciones["E3:E1000"].formula1 == "=CatalogoASIC"
    assert validaciones["C3:C1000"].formula1 == "=CatalogoMunicipios"
    assert validaciones["D3:D1000"].formula1 == (
        '=INDIRECT("Parroquia"&MATCH(C3,CatalogoMunicipios,0))'
    )
    assert workbook["Catalogos"].sheet_state == "visible"
    assert validaciones["E3:E1000"].showDropDown is False
    assert validaciones["E3:E1000"].showInputMessage is True
    assert validaciones["C3:C1000"].showDropDown is False
    assert validaciones["D3:D1000"].showDropDown is False
    assert workbook.defined_names["CatalogoASIC"].attr_text == (
        "'Catalogos'!$G$2:$G$26"
    )
    assert workbook.defined_names["CatalogoMunicipios"].attr_text == (
        "'Catalogos'!$H$2:$H$30"
    )
    assert workbook.defined_names["Parroquia1"].attr_text == (
        "'Catalogos'!$I$2:$I$2"
    )
    assert workbook.defined_names["Parroquia29"].attr_text == (
        "'Catalogos'!$AK$2:$AK$5"
    )
    assert [
        workbook["Catalogos"].cell(row=fila, column=7).value
        for fila in range(2, 27)
    ] == sorted({
        ubicacion["asic_nombre"] for ubicacion in _ubicaciones()
    })
    assert all(not nombre.startswith("$") for nombre in workbook.defined_names)


def test_plantillas_distribuidas_tienen_listas_dependientes():
    actividad = load_workbook("assets/templates/plantilla_importacion_actividad.xlsx")
    hoja_actividad = actividad.active
    validaciones_actividad = {
        str(validacion.sqref): validacion
        for validacion in hoja_actividad.data_validations.dataValidation
    }
    assert validaciones_actividad["E4:E1048576"].formula1 == "=CatalogoASIC"
    assert validaciones_actividad["C4:C1048576"].formula1 == "=CatalogoMunicipios"
    assert validaciones_actividad["D4:D1048576"].formula1 == (
        '=INDIRECT("Parroquia"&MATCH(C4,CatalogoMunicipios,0))'
    )
    assert actividad["Catalogos"].sheet_state == "visible"

    vocerias = load_workbook("assets/templates/plantilla_importacion_voceria.xlsx")
    hoja_vocerias = vocerias.active
    validaciones_vocerias = {
        str(validacion.sqref): validacion
        for validacion in hoja_vocerias.data_validations.dataValidation
    }
    assert validaciones_vocerias["E3:E1000"].formula1 == "=CatalogoASIC"
    assert vocerias.defined_names["CatalogoASIC"].attr_text == "'Catalogos'!$G$2:$G$26"
    assert validaciones_vocerias["C3:C1000"].formula1 == "=CatalogoMunicipios"
    assert validaciones_vocerias["D3:D1000"].formula1 == (
        '=INDIRECT("Parroquia"&MATCH(C3,CatalogoMunicipios,0))'
    )
    assert vocerias["Catalogos"].sheet_state == "visible"


def test_importacion_de_actividades_rechaza_parroquia_ajena_al_municipio():
    class _ParroquiaDAO:
        @staticmethod
        def listar_por_asic(id_asic):
            return [{"municipio": "San Cristóbal", "parroquia": "San Sebastián"}]

    service = ExcelService.__new__(ExcelService)
    service._parroquia_dao = _ParroquiaDAO()
    service._cache_ubicaciones = {}
    fila = [
        "Táchira",
        "San Cristóbal",
        "San Pablo",
        ACTIVIDAD_ESTABLECIMIENTO_OPCIONES[0],
        "Centro de Salud",
    ]
    idx = {
        "estado": 0,
        "municipio": 1,
        "parroquia": 2,
        "establecimiento_salud": 3,
        "nombre_establecimiento_salud": 4,
    }

    errores = service._validar_campos_fila(
        fila, idx, date.today(), "Tema", 0, 0, 20
    )

    assert "'PARROQUIA' no corresponde al municipio seleccionado." in errores


def test_importacion_de_vocerias_rechaza_parroquia_ajena_al_municipio():
    class _AsicDAO:
        @staticmethod
        def obtener_todos():
            return [{"id_asic": 20, "nombre": "Sabaneta"}]

    class _ParroquiaDAO:
        @staticmethod
        def listar_por_asic(id_asic):
            return [{"municipio": "San Cristóbal", "parroquia": "San Sebastián"}]

    service = ExcelServiceVoceria.__new__(ExcelServiceVoceria)
    service._asic_dao = _AsicDAO()
    service._parroquia_dao = _ParroquiaDAO()
    service._cache_asic = {}
    service._cache_ubicaciones = {}
    datos = {
        "estado": "Táchira",
        "municipio": "San Cristóbal",
        "parroquia": "San Pablo",
        "asic": "Sabaneta",
        "nombre_comite": "Comité de prueba",
        "centro_dispensador": "CDI",
        "nombre_centro": "Centro de Salud",
        "nombres": "Nombre Apellido",
        "ci": "1234567",
        "telefono": "04121112222",
        "voceria": "V. CONTRALORÍA",
        "fecha_nacimiento": "2000-01-01",
        "edad": str(date.today().year - 2000),
        "sexo": "M",
        "profesion": "Médico",
        "correo": "persona@example.com",
        "fechas_eleccion": date.today().strftime("%Y-%m-%d"),
    }

    errores = service._validar_fila(datos, 3)

    assert any(
        "La parroquia no corresponde al municipio seleccionado" in error
        for error in errores
    )


def test_guardado_de_actividad_rechaza_parroquia_ajena_al_municipio():
    class _ParroquiaDAO:
        @staticmethod
        def listar_por_asic(id_asic):
            return [{"municipio": "San Cristóbal", "parroquia": "San Sebastián"}]

    controller = ActividadController.__new__(ActividadController)
    controller.parroquia_dao = _ParroquiaDAO()

    resultado = controller.guardar_actividad(
        id_asic=20,
        id_tipo=1,
        fecha=date.today().isoformat(),
        estado="Táchira",
        municipio="San Cristóbal",
        parroquia="San Pablo",
        establecimiento_salud=ACTIVIDAD_ESTABLECIMIENTO_OPCIONES[0],
        nombre_establecimiento="Centro de Salud",
        comites_salud=0,
    )

    assert resultado == (
        False,
        "La parroquia debe corresponder al municipio seleccionado.",
    )


def test_guardado_de_voceria_rechaza_parroquia_ajena_al_municipio():
    class _ParroquiaDAO:
        @staticmethod
        def listar_por_asic(id_asic):
            return [{"municipio": "San Cristóbal", "parroquia": "San Sebastián"}]

    controller = VoceriaController.__new__(VoceriaController)
    controller.parroquia_dao = _ParroquiaDAO()

    resultado = controller.guardar_voceria(
        id_asic=20,
        estado="Táchira",
        municipio="San Cristóbal",
        parroquia="San Pablo",
        centro_dispensador=VOCERIA_CENTRO_DISPENSADOR_OPCIONES[0],
        nombre_centro="Centro de Salud",
        nombre_comite="Comité de prueba",
        items=[],
    )

    assert resultado == (
        False,
        "La parroquia debe corresponder al municipio seleccionado.",
    )
