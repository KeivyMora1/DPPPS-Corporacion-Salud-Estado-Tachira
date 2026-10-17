import os
import threading
from datetime import date, datetime
from tkinter import filedialog, messagebox

import customtkinter as ctk
from tkinter import ttk

from constants import ( TEMPLATE_PATH )
from db import ExcelService, AuthController
from utils import set_favicon

from .model import ModeloActividad
from .view import VistaActividad

def generar_plantilla_actividad(
    ruta_entrada, ruta_salida, nombres_tipos, ubicaciones_asic=None
):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.utils import get_column_letter

    if ruta_entrada == ruta_salida:
        raise ValueError("La plantilla generada debe guardarse en un archivo distinto al original.")
    if not nombres_tipos:
        raise ValueError("No hay tipos de actividad disponibles para incluir en la plantilla.")
    if ubicaciones_asic is None:
        from db import ParroquiaDAO
        ubicaciones_asic = ParroquiaDAO().obtener_todas()
    if not ubicaciones_asic:
        raise ValueError("No hay ubicaciones ASIC disponibles para incluir en la plantilla.")

    workbook = openpyxl.load_workbook(ruta_entrada)
    hoja_actividad = workbook.active

    for nombre_hoja in ("_ubicaciones", "Catalogos"):
        if nombre_hoja in workbook.sheetnames:
            del workbook[nombre_hoja]
            
    hoja_ubicaciones = workbook.create_sheet("Catalogos")

    nombres_asics = sorted({
        ubicacion["asic_nombre"] for ubicacion in ubicaciones_asic
    })

    parroquias_por_municipio = {}
    for ubicacion in ubicaciones_asic:
        municipio = ubicacion["municipio"]
        parroquias_por_municipio.setdefault(municipio, set()).add(
            ubicacion["parroquia"]
        )

    municipios = sorted(parroquias_por_municipio)

    def agregar_lista_dependiente(nombre, columna, opciones):
        for fila, opcion in enumerate(opciones, start=2):
            hoja_ubicaciones.cell(row=fila, column=columna, value=opcion)
        letra = get_column_letter(columna)
        hoja_ubicaciones.cell(row=1, column=columna, value=nombre)
        
        if nombre in workbook.defined_names:
            del workbook.defined_names[nombre]
            
        workbook.defined_names.add(DefinedName(
            nombre,
            attr_text=(
                f"'{hoja_ubicaciones.title}'!"
                f"${letra}$2:${letra}${len(opciones) + 1}"
            ),
        ))

    agregar_lista_dependiente("CatalogoASIC", 1, nombres_asics)
    agregar_lista_dependiente("CatalogoMunicipios", 2, municipios)

    columna_parroquias = 3
    for indice, municipio in enumerate(municipios, start=1):
        col_idx = columna_parroquias + indice - 1
        agregar_lista_dependiente(
            f"Parroquia{indice}", col_idx,
            sorted(parroquias_por_municipio[municipio]),
        )
        hoja_ubicaciones.cell(row=1, column=col_idx, value=municipio)

    max_parroquia_idx = len(municipios)
    for nombre in list(workbook.defined_names):
        if nombre.startswith("Parroquia"):
            try:
                idx = int(nombre.replace("Parroquia", ""))
                if idx > max_parroquia_idx:
                    del workbook.defined_names[nombre]
            except ValueError:
                pass

    rangos_validacion = {
        "E4:E1048576": "=CatalogoASIC",
        "C4:C1048576": "=CatalogoMunicipios",
        "D4:D1048576": '=INDIRECT("Parroquia"&MATCH(C4,CatalogoMunicipios,0))',
    }
    for validacion in list(hoja_actividad.data_validations.dataValidation):
        if str(validacion.sqref) in {*rangos_validacion, "I4:I1048576"}:
            hoja_actividad.data_validations.dataValidation.remove(validacion)

    for rango, formula in rangos_validacion.items():
        validacion = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            showDropDown=False,
        )
        validacion.add(rango)
        hoja_actividad.add_data_validation(validacion)

    opciones_directas = ",".join(nombres_tipos)
    if len(opciones_directas) <= 253:
        formula_tipos = f'"{opciones_directas}"'
    else:
        columna_opciones = hoja_actividad.max_column + 1
        letra_opciones = get_column_letter(columna_opciones)
        for fila, nombre_tipo in enumerate(nombres_tipos, start=1):
            hoja_actividad.cell(row=fila, column=columna_opciones, value=nombre_tipo)
        hoja_actividad.column_dimensions[letra_opciones].hidden = True
        formula_tipos = f"=${letra_opciones}$1:${letra_opciones}${len(nombres_tipos)}"

    validacion_tipos = DataValidation(
        type="list",
        formula1=formula_tipos,
        allow_blank=True,
        showErrorMessage=True,
        showDropDown=False,
    )
    validacion_tipos.add("I4:I1048576")
    hoja_actividad.add_data_validation(validacion_tipos)

    workbook.save(ruta_salida)


class _ImportarModal:
    def __init__(self, parent_view):
        self.parent_view = parent_view
        self.parent = parent_view.parent
        self.t = parent_view.t
        self.controller = parent_view.controlador.modelo.controller
        self.auth = self.controller.actividad_dao.auth if hasattr(self.controller.actividad_dao, 'auth') else None
        self.auth_ctrl = AuthController()
        self.usuario = self.auth_ctrl.obtener_usuario_actual()
        self._ruta_archivo = None
        self._resultado_previo = None

        self._construir_modal()
        self._mostrar_paso1()

    def _construir_modal(self):
        t = self.t
        self.ventana = ctk.CTkToplevel(self.parent)
        set_favicon(self.ventana)
        self.ventana.title("Importar Actividades desde Excel")
        self.ventana.geometry("850x700")
        self.ventana.resizable(False, False)
        self.ventana.configure(fg_color=t["bg_surface"])
        self.ventana.grab_set()

        self.ventana.update_idletasks()
        x = (self.ventana.winfo_screenwidth() // 2) - 425
        y = (self.ventana.winfo_screenheight() // 2) - 350
        self.ventana.geometry(f"+{x}+{y}")

        header = ctk.CTkFrame(self.ventana, fg_color="transparent")
        header.pack(fill="x", padx=20, pady=(15, 5))

        ctk.CTkLabel(header, text="📥 Importar Actividades desde Excel",
                     font=("Arial", 17, "bold"), text_color=t["text_primary"]).pack(side="left")

        ctk.CTkButton(header, text="✕", width=30, height=30,
                      fg_color="transparent", hover_color=t["bg_hover"],
                      text_color=t["text_primary"], corner_radius=6,
                      command=self.ventana.destroy).pack(side="right")

        ctk.CTkLabel(self.ventana,
                     text="Importa actividades en bloque subiendo un archivo .xlsx con el formato de la plantilla.",
                     font=("Arial", 11), text_color=t["text_muted"],
                     ).pack(anchor="w", padx=20, pady=(0, 10))

        self._frame_pasos = ctk.CTkFrame(self.ventana, fg_color="transparent")
        self._frame_pasos.pack(fill="both", expand=True, padx=20, pady=(0, 15))

        btn_frame = ctk.CTkFrame(self.ventana, fg_color="transparent")
        btn_frame.pack(fill="x", padx=20, pady=(0, 15))

        self._btn_cerrar = ctk.CTkButton(btn_frame, text="✕  Cerrar",
                         width=118, height=38,
                         font=("Arial", 12, "bold"), fg_color=t["bg_card"],
                         hover_color=t["bg_hover"], text_color=t["text_primary"],
                         border_width=1, border_color=t["accent"], corner_radius=8,
                                         command=self.ventana.destroy)
        self._btn_cerrar.pack(side="right")

    def _limpiar_pasos(self):
        for w in self._frame_pasos.winfo_children():
            w.destroy()

    def _mostrar_paso1(self):
        t = self.t
        self._limpiar_pasos()
        self._ruta_archivo = None
        self._resultado_previo = None

        card = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card.pack(fill="x", padx=10)

        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="x", padx=20, pady=20)

        ctk.CTkLabel(inner, text="Paso 1 de 3 — Seleccionar archivo Excel",
                     font=("Arial", 14, "bold"), text_color=t["text_primary"]).pack(anchor="w", pady=(0, 6))

        ctk.CTkLabel(inner,
                     text="El archivo debe seguir la plantilla descargable y contener al menos: FECHA, ASIC, ACTIVIDAD DESARROLLADA, TEMA TRATADO, LUGAR/UBICACIÓN y TOTAL DE ASISTENCIA.",
                     font=("Arial", 10), text_color=t["text_muted"],
                     wraplength=750, justify="left").pack(anchor="w", pady=(0, 6))

        ctk.CTkLabel(inner,
                     text="CANTIDAD DE COMITÉS DE SALUD debe ser un número entero de 0 a 99, no el nombre del comité.",
                     font=("Arial", 10), text_color=t["text_muted"],
                     wraplength=650, justify="left").pack(anchor="w", pady=(0, 15))

        self._lbl_ruta = ctk.CTkLabel(inner, text="Ningún archivo seleccionado",
                                      font=("Arial", 11, "italic"), text_color=t["text_muted"])
        self._lbl_ruta.pack(anchor="w", pady=(0, 12))

        fila_botones = ctk.CTkFrame(inner, fg_color="transparent")
        fila_botones.pack(anchor="w")

        ctk.CTkButton(fila_botones, text="📂 Examinar…",
                      width=140, height=36, font=("Arial", 11), fg_color=t["bg_card"],
                      hover_color=t["bg_hover"], text_color=t["text_primary"],
                      corner_radius=8, command=self._seleccionar_archivo).pack(side="left", padx=(0, 10))

        self._btn_previsualizar = ctk.CTkButton(
            fila_botones, text="▶ Previsualizar",
            width=150, height=36, font=("Arial", 11, "bold"), fg_color=t["bg_hover"],
            hover_color=t["bg_hover"], text_color=t["text_dim"],
            text_color_disabled=t["text_dim"], corner_radius=8,
            state="disabled", command=self._iniciar_previo)
        self._btn_previsualizar.pack(side="left")

    def _seleccionar_archivo(self):
        t = self.t
        ruta = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=[("Archivos Excel compatibles", "*.xlsx")],
            parent=self.ventana
        )
        if ruta:
            self._ruta_archivo = ruta
            nombre = os.path.basename(ruta)
            self._lbl_ruta.configure(text=f"📄 {nombre}  ({ruta})", text_color=t["text_primary"])
            self._btn_previsualizar.configure(
                state="normal", fg_color=t["accent"],
                hover_color=t["accent_dim"], text_color=t["text_primary"])

    def _iniciar_previo(self):
        self._btn_previsualizar.configure(state="disabled", text="Procesando…")
        self.ventana.configure(cursor="wait")

        def _tarea():
            try:
                svc = ExcelService(id_usuario_registro=self.usuario["id_usuario"])
                resultado = svc.previsualizar(self._ruta_archivo)
                self.ventana.after(0, lambda: self._mostrar_paso2(resultado))
            except Exception as exc:
                self.ventana.after(0, lambda e=exc: self._error_previo(e))

        threading.Thread(target=_tarea, daemon=True).start()

    def _error_previo(self, exc):
        self.ventana.configure(cursor="")
        messagebox.showerror("Error al leer el archivo",
                             f"No se pudo procesar el archivo:\n\n{exc}", parent=self.ventana)
        self._btn_previsualizar.configure(state="normal", text="▶ Previsualizar",
                                          fg_color=self.t["bg_hover"],
                                          text_color=self.t["text_dim"])

    def _mostrar_paso2(self, resultado):
        t = self.t
        self.ventana.configure(cursor="")
        self._resultado_previo = resultado
        self._limpiar_pasos()

        validas = resultado["filas_validas"]
        errores = resultado["filas_error"]
        total = resultado["total_filas"]

        card_resumen = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card_resumen.pack(fill="x", pady=(0, 10), padx=10)

        inner = ctk.CTkFrame(card_resumen, fg_color="transparent")
        inner.pack(fill="x", padx=15, pady=12)

        ctk.CTkLabel(inner, text="Paso 2 de 3 — Resultado del análisis",
                     font=("Arial", 14, "bold"), text_color=t["text_primary"]).pack(anchor="w", pady=(0, 10))

        fila_resumen = ctk.CTkFrame(inner, fg_color="transparent")
        fila_resumen.pack(anchor="w")

        def _chip(parent, texto, color):
            chip = ctk.CTkFrame(parent, fg_color=color, corner_radius=6)
            chip.pack(side="left", padx=(0, 8))
            ctk.CTkLabel(chip, text=texto, font=("Arial", 10, "bold"),
                         text_color=t["text_primary"]).pack(padx=10, pady=5)

        _chip(fila_resumen, f"📄 Total filas: {total}", t["bg_card"])
        _chip(fila_resumen, f"✅ Válidas: {len(validas)}", t["success"])
        _chip(fila_resumen, f"❌ Con errores: {len(errores)}",
              t["red"] if errores else t["bg_hover"])

        if validas:
            titulo_validas = "Filas listas para importar" if not errores else "Filas sin errores individuales"
            self._seccion_tabla(titulo_validas, t["success"], validas, "validas")
        if errores:
            self._seccion_tabla("Filas con errores (no se importarán)", t["red"], errores, "errores")

        card_acciones = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card_acciones.pack(fill="x", padx=10, pady=(8, 0))

        inner_acc = ctk.CTkFrame(card_acciones, fg_color="transparent")
        inner_acc.pack(fill="x", padx=15, pady=12)

        ctk.CTkButton(inner_acc, text="← Seleccionar otro archivo",
                      height=38, font=("Arial", 11), fg_color=t["bg_card"],
                      hover_color=t["bg_hover"], text_color=t["text_primary"],
                      corner_radius=8, command=self._mostrar_paso1).pack(side="left", padx=(0, 10))

        if validas and not errores:
            ctk.CTkButton(inner_acc,
                          text=f"✔ Importar {len(validas)} actividades",
                          height=38, font=("Arial", 11, "bold"), fg_color=t["success"],
                          hover_color=t["success_dim"], text_color=t["text_primary"],
                          corner_radius=8, command=self._confirmar_importar).pack(side="left")
        else:
            ctk.CTkLabel(inner_acc,
                         text=("El archivo contiene errores; no se importará ninguna actividad. Corrige el archivo y vuelve a intentarlo."
                               if errores else "No hay filas válidas para importar. Corrige el archivo y vuelve a intentarlo."),
                         font=("Arial", 10, "italic"), text_color=t["red"]).pack(side="left")

    def _seccion_tabla(self, titulo, color_titulo, filas, modo):
        t = self.t
        marco = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        marco.pack(fill="both", expand=(modo == "validas"), pady=(0, 8), padx=10)

        inner = ctk.CTkFrame(marco, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=10, pady=8)

        ctk.CTkLabel(inner, text=titulo, font=("Arial", 11, "bold"),
                     text_color=color_titulo).pack(anchor="w", pady=(0, 4))

        if modo == "validas":
            columnas = ("fila", "asic", "tipo", "fecha", "participantes")
        else:
            columnas = ("fila", "asic", "tipo", "fecha", "errores")

        tree_frame = ctk.CTkFrame(inner, fg_color="transparent")
        tree_frame.pack(fill="both", expand=True)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("ImportModal.Treeview",
                        background=t["bg_base"], foreground=t["text_primary"],
                        fieldbackground=t["bg_base"], rowheight=30,
                        font=("Arial", 11))
        style.configure("ImportModal.Treeview.Heading",
                        background=t["bg_card"], foreground=t["text_primary"],
                        font=("Arial", 11, "bold"), relief="flat")
        style.map("ImportModal.Treeview",
                  background=[("selected", t["accent"])],
                  foreground=[("selected", t["text_primary"])])

        tree = ttk.Treeview(tree_frame, columns=columnas, show="headings",
                            style="ImportModal.Treeview", height=min(len(filas), 8))

        col_ancho = {"fila": 40, "asic": 150, "tipo": 130, "fecha": 90, "participantes": 80, "errores": 320}
        if modo == "errores":
            col_ancho["errores"] = 850
        col_labels = {"fila": "#", "asic": "ASIC", "tipo": "Tipo Actividad",
                      "fecha": "Fecha", "participantes": "Partic.", "errores": "Errores detectados"}

        for col in columnas:
            tree.heading(col, text=col_labels[col], anchor="center")
            tree.column(col, width=col_ancho[col], minwidth=40,
                        anchor="w" if col == "errores" else "center",
                        stretch=modo != "errores")

        def fecha_sin_hora(valor):
            if isinstance(valor, datetime):
                return valor.strftime("%Y-%m-%d")
            if isinstance(valor, date):
                return valor.isoformat()
            return valor or ""

        for fila in filas:
            dc = fila.get("datos_crudos", {})
            if modo == "validas":
                valores = (
                    fila["numero_fila"],
                    dc.get("asic", ""),
                    dc.get("tipo_actividad", ""),
                    fecha_sin_hora(fila.get("fecha") or dc.get("fecha")),
                    fila["participantes"],
                )
            else:
                valores = (
                    fila["numero_fila"],
                    dc.get("asic", ""),
                    dc.get("tipo_actividad", ""),
                    fecha_sin_hora(dc.get("fecha")),
                    " | ".join(fila.get("errores", [])),
                )
            tree.insert("", "end", values=valores)

        scrolly = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        scrollx = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=scrolly.set, xscrollcommand=scrollx.set)
        scrolly.pack(side="right", fill="y")
        scrollx.pack(side="bottom", fill="x")
        tree.pack(fill="both", expand=True)

        tree.bind("<Button-1>", self._on_tabla_click, add="+")

    def _on_tabla_click(self, event):
        tree = event.widget
        item = tree.identify_row(event.y)
        if not item:
            tree.selection_remove(tree.selection())

    def _confirmar_importar(self):
        if self._resultado_previo.get("filas_error"):
            messagebox.showwarning(
                "Archivo con errores",
                "Corrige todas las filas con errores antes de importar. No se guardó ninguna actividad.",
                parent=self.ventana,
            )
            return
        validas = self._resultado_previo["filas_validas"]
        if messagebox.askyesno("Confirmar importación",
                               f"Se importarán {len(validas)} actividades en la base de datos.\n\n¿Deseas continuar?",
                               parent=self.ventana):
            self._ejecutar_importacion(validas)

    def _ejecutar_importacion(self, filas_validas):
        self.ventana.configure(cursor="wait")

        def _tarea():
            try:
                svc = ExcelService(id_usuario_registro=self.usuario["id_usuario"])
                resultado = svc.importar(filas_validas)
                self.ventana.after(0, lambda r=resultado: self._mostrar_paso3(r))
            except Exception as exc:
                self.ventana.after(0, lambda e=exc: self._error_importar(e))

        threading.Thread(target=_tarea, daemon=True).start()

    def _error_importar(self, exc):
        self.ventana.configure(cursor="")
        messagebox.showerror("Error al importar",
                             f"Ocurrió un error durante la importación:\n\n{exc}", parent=self.ventana)

    def _mostrar_paso3(self, resultado):
        t = self.t
        self.ventana.configure(cursor="")
        self._limpiar_pasos()

        insertadas = resultado["insertadas"]
        errores_bd = resultado["errores_al_insertar"]

        card = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card.pack(fill="x", padx=10)

        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="x", padx=25, pady=25)

        icono = "✅" if not errores_bd else "⚠️"
        color = t["success"] if not errores_bd else t["amber"]
        titulo = "Importación completada" if not errores_bd else "Importación completada con advertencias"

        ctk.CTkLabel(inner, text=f"{icono}  {titulo}",
                     font=("Arial", 16, "bold"), text_color=color).pack(anchor="w", pady=(0, 12))

        ctk.CTkLabel(inner, text=f"Se insertaron correctamente {insertadas} actividades.",
                     font=("Arial", 12), text_color=t["text_primary"]).pack(anchor="w")

        if errores_bd:
            ctk.CTkLabel(inner,
                         text=f"{len(errores_bd)} actividades no pudieron insertarse (ver detalles abajo).",
                         font=("Arial", 11), text_color=t["red"]).pack(anchor="w", pady=(4, 0))
            for err in errores_bd:
                ctk.CTkLabel(inner, text=f"  • Fila {err['numero_fila']}: {err['error']}",
                             font=("Arial", 10), text_color=t["red"]).pack(anchor="w")

        ctk.CTkButton(inner, text="↩ Hacer otra importación",
                      height=38, font=("Arial", 11), fg_color=t["bg_card"],
                      hover_color=t["bg_hover"], text_color=t["text_primary"],
                      corner_radius=8, command=self._mostrar_paso1).pack(anchor="w", pady=(18, 0))

        ctk.CTkButton(inner, text="✕  Cerrar",
                  width=118, height=38, font=("Arial", 12, "bold"), fg_color=t["accent"],
                      hover_color=t["accent_dim"], text_color=t["text_primary"],
                      corner_radius=8, command=lambda: (self.ventana.destroy(), self.parent_view.controlador.cargar_datos_tabla())
                      ).pack(anchor="w", pady=(8, 0))

class ControladorActividad:
    def __init__(self, parent):
        self.modelo = ModeloActividad()

        self.asic_map = {}
        self.ubicaciones_por_asic = {}
        self.tipo_map = {}
        self.filtros_actuales = None
        self._cargar_catalogos()

        self.vista = VistaActividad(parent, self)
        self.cargar_datos_tabla()

    def puede_escribir(self):
        return self.modelo.controller.auth.puede_escribir()

    def _verificar_permiso_escritura(self):
        if self.puede_escribir():
            return True
        messagebox.showwarning("Permiso denegado", "Tu rol de lector no permite esta acción.")
        return False

    def _cargar_catalogos(self):
        self.ubicaciones_por_municipio = {}
        try:
            asics = self.modelo.listar_asics()
            for ubicacion in self.modelo.listar_ubicaciones_asic():
                self.ubicaciones_por_asic.setdefault(
                    ubicacion["asic_nombre"], []
                ).append(ubicacion)
                self.ubicaciones_por_municipio.setdefault(
                    ubicacion["municipio"], []
                ).append(ubicacion)
            self.asic_map = {
                asic["nombre"]: asic["id_asic"]
                for asic in asics
                if asic["nombre"] in self.ubicaciones_por_asic
            }

            tipos = self.modelo.listar_tipos_actividad()
            self.tipo_map = {t["nombre"]: t["id_tipo"] for t in tipos}
        except Exception as e:
            messagebox.showerror("Error de base de datos", f"No se pudieron cargar los catalogos: {str(e)}")

    def aplicar_filtros(self):
        asic_sel = self.vista.combo_filtro_asic.get()
        tipo_sel = self.vista.combo_filtro_tipo.get()
        f_desde = self.vista.ent_filtro_desde.get().strip()
        f_hasta = self.vista.ent_filtro_hasta.get().strip()

        filtros = {}

        if asic_sel != "Todos":
            filtros["id_asic"] = self.asic_map.get(asic_sel)
        if tipo_sel != "Todos":
            filtros["id_tipo"] = self.tipo_map.get(tipo_sel)

        if f_desde and f_desde != "AAAA-MM-DD":
            try:
                datetime.strptime(f_desde, "%Y-%m-%d")
                filtros["fecha_desde"] = f_desde
            except ValueError:
                messagebox.showerror("Error en Filtros", "La fecha 'Desde' debe tener el formato AAAA-MM-DD (ej: 2026-05-17).")
                return

        if f_hasta and f_hasta != "AAAA-MM-DD":
            try:
                datetime.strptime(f_hasta, "%Y-%m-%d")
                filtros["fecha_hasta"] = f_hasta
            except ValueError:
                messagebox.showerror("Error en Filtros", "La fecha 'Hasta' debe tener el formato AAAA-MM-DD (ej: 2026-05-17).")
                return

        if filtros.get("fecha_desde") and filtros.get("fecha_hasta") and filtros["fecha_desde"] > filtros["fecha_hasta"]:
            messagebox.showerror("Error en Filtros", "La fecha 'Desde' no puede ser posterior a 'Hasta'.")
            return

        self.filtros_actuales = filtros
        self.cargar_datos_tabla()

    def limpiar_filtros(self):
        self.vista.combo_filtro_asic.set("Todos")
        self.vista.combo_filtro_tipo.set("Todos")
        self.vista.ent_filtro_desde.set_date("")
        self.vista.ent_filtro_hasta.set_date("")
        self.filtros_actuales = None
        self.cargar_datos_tabla()

    def cargar_datos_tabla(self, filtros=None):
        for item in self.vista.tabla.get_children():
            self.vista.tabla.delete(item)

        try:
            if filtros is not None:
                self.filtros_actuales = filtros
            lista = self.modelo.listar_actividades(self.filtros_actuales)
            for act in lista:
                f = act["fecha"]
                if isinstance(f, (date, datetime)):
                    f = f.strftime("%d/%m/%Y")

                self.vista.tabla.insert("", "end", iid=str(act["id_actividad"]), values=(
                    act["asic_nombre"],
                    act["tipo_nombre"],
                    f,
                    act.get("estado") or "",
                    act.get("municipio") or "",
                    act.get("parroquia") or "",
                    act.get("establecimiento_salud") or "",
                    act.get("nombre_establecimiento") or "",
                    act.get("comites_salud") if act.get("comites_salud") is not None else "",
                    act.get("total_participantes", act.get("participantes") or 0),
                ))
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrio un error al cargar las actividades: {str(e)}")

    def abrir_form_nueva(self):
        if not self._verificar_permiso_escritura():
            return
        self.vista.abrir_formulario()

    def editar_seleccionada(self):
        if not self._verificar_permiso_escritura():
            return
        seleccion = self.vista.tabla.selection()
        if not seleccion:
            messagebox.showwarning("Atención", "Selecciona una actividad de la tabla para editar.")
            return
        try:
            id_actividad = int(seleccion[0])
            act = self.modelo.obtener_actividad_por_id(id_actividad)
            if not act:
                messagebox.showerror("Error", "No se pudo recuperar la información de la actividad.")
                return
            self.vista.abrir_formulario(act)
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al obtener la actividad: {str(e)}")

    def eliminar_seleccionada(self):
        if not self._verificar_permiso_escritura():
            return
        seleccion = self.vista.tabla.selection()
        if not seleccion:
            messagebox.showwarning("Atención", "Selecciona una actividad de la tabla para eliminar.")
            return
        try:
            id_actividad = int(seleccion[0])
            valores = self.vista.tabla.item(seleccion[0])["values"]
            asic_act = valores[0] if len(valores) > 0 else ""
            tipo_act = valores[1] if len(valores) > 1 else ""
            fecha_act = valores[2] if len(valores) > 2 else ""

            if messagebox.askyesno("Confirmar eliminación",
                                   f"¿Deseas eliminar la actividad '{tipo_act}' del día {fecha_act} ({asic_act})?\nEsta acción no se puede deshacer."):
                ok, msg = self.modelo.eliminar_actividad(id_actividad)
                if ok:
                    messagebox.showinfo("Éxito", msg)
                    self.cargar_datos_tabla()
                else:
                    messagebox.showerror("Error", msg)
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo eliminar la actividad: {str(e)}")

    def on_seleccion_cambio(self, event=None):
        seleccion = self.vista.tabla.selection()
        if seleccion:
            self.vista._mostrar_botones_accion()
        else:
            self.vista._ocultar_botones_accion()

    def descargar_plantilla(self):
        if not os.path.exists(TEMPLATE_PATH):
            messagebox.showerror("Plantilla no encontrada",
                                 f"No se encontro la plantilla en:\n{TEMPLATE_PATH}\n\n"
                                 "Asegurate de que el archivo exista en esa ruta.")
            return

        destino = filedialog.asksaveasfilename(
            title="Guardar plantilla como...",
            defaultextension=".xlsx",
            initialfile="plantilla_importacion_actividad.xlsx",
            filetypes=[("Archivos Excel", "*.xlsx")],
        )
        if not destino:
            return

        try:
            nombres_tipos = [
                tipo["nombre"] for tipo in self.modelo.listar_tipos_actividad()
            ]
            generar_plantilla_actividad(
                TEMPLATE_PATH,
                destino,
                nombres_tipos,
                self.modelo.listar_ubicaciones_asic(),
            )
        except Exception as exc:
            messagebox.showerror(
                "Error al generar la plantilla",
                f"No se pudo generar la plantilla de importación: {exc}",
            )
            return

        messagebox.showinfo("Plantilla guardada", f"Plantilla guardada en:\n{destino}")

    def exportar_consolidado(self):
        import openpyxl
        from copy import copy

        try:
            actividades = self.modelo.listar_actividades(self.filtros_actuales)
            if not actividades:
                messagebox.showinfo("Sin datos", "No hay actividades para exportar con los filtros actuales.")
                return

            destino = filedialog.asksaveasfilename(
                title="Guardar consolidado de actividades como...",
                defaultextension=".xlsx",
                initialfile="Consolidado_Actividades.xlsx",
                filetypes=[("Archivos Excel", "*.xlsx")],
            )
            if not destino:
                return

            if not os.path.exists(TEMPLATE_PATH):
                messagebox.showerror("Plantilla no encontrada", f"No se encuentra la plantilla en:\n{TEMPLATE_PATH}")
                return

            workbook = openpyxl.load_workbook(TEMPLATE_PATH)
            sheet = workbook.active
            fila_modelo = 4

            def copiar_estilo_fila(fila_destino):
                for columna in range(1, sheet.max_column + 1):
                    celda_modelo = sheet.cell(row=fila_modelo, column=columna)
                    celda_destino = sheet.cell(row=fila_destino, column=columna)
                    if celda_modelo.has_style:
                        celda_destino._style = copy(celda_modelo._style)
                    if celda_modelo.number_format:
                        celda_destino.number_format = celda_modelo.number_format
                sheet.row_dimensions[fila_destino].height = sheet.row_dimensions[fila_modelo].height

            campos_participantes = (
                "cant_medicos",
                "cant_promotores_salud",
                "cant_voceros_comite_salud",
                "cant_vocero_consejo_direccion",
                "cant_coordinador_docencia",
                "cant_autoridades",
                "cant_otros_actores",
                "cant_personas_comunidad",
                "cant_estudiantes_aspirantes",
            )

            for fila, actividad in enumerate(actividades, start=fila_modelo):
                if fila != fila_modelo:
                    copiar_estilo_fila(fila)

                fecha_actividad = actividad.get("fecha")
                if isinstance(fecha_actividad, datetime):
                    fecha_actividad = fecha_actividad.date()

                valores = (
                    fecha_actividad or "",
                    actividad.get("estado") or "",
                    actividad.get("municipio") or "",
                    actividad.get("parroquia") or "",
                    actividad.get("asic_nombre") or "",
                    actividad.get("establecimiento_salud") or "",
                    actividad.get("nombre_establecimiento") or "",
                    actividad.get("comites_salud") if actividad.get("comites_salud") is not None else "",
                    actividad.get("tipo_nombre") or "",
                    actividad.get("tema") or "",
                    *(actividad.get(campo) if actividad.get(campo) is not None else "" for campo in campos_participantes),
                    actividad.get("total_participantes") if actividad.get("total_participantes") is not None else "",
                )
                for columna, valor in enumerate(valores, start=1):
                    sheet.cell(row=fila, column=columna, value=valor)
                sheet.cell(row=fila, column=1).number_format = "dd-mm-yyyy"

            workbook.save(destino)
            messagebox.showinfo("Éxito", f"Consolidado exportado correctamente a:\n{destino}")
        except Exception as exc:
            messagebox.showerror("Error", f"Ocurrió un error al exportar: {str(exc)}")

    def abrir_modal_importar(self):
        if not self._verificar_permiso_escritura():
            return
        _ImportarModal(self.vista)


__all__ = ['ControladorActividad', 'generar_plantilla_actividad']
