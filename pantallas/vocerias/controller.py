import os
import re
import shutil
import unicodedata
import threading
from datetime import date, datetime

import customtkinter as ctk
from tkinter import filedialog, messagebox
from tkinter import ttk

from db import VOCERIA_TEMPLATE_PATH, build_voceria_template, ExcelServiceVoceria
from .model import ModeloVoceria
from .view import VistaVoceria


class _ImportarVoceriaModal:
    def __init__(self, parent_view):
        self.parent_view = parent_view
        self.parent = parent_view.parent
        self.t = parent_view.t
        self.controller = parent_view.controlador.modelo.controller
        from db import AuthController
        self.auth_ctrl = AuthController()
        self.usuario = self.auth_ctrl.obtener_usuario_actual()
        self._ruta_archivo = None
        self._resultado_previo = None

        self._construir_modal()
        self._mostrar_paso1()

    def _construir_modal(self):
        t = self.t
        self.ventana = ctk.CTkToplevel(self.parent)
        self.ventana.title("Importar Vocerías desde Excel")
        self.ventana.geometry("850x700")
        self.ventana.resizable(False, False)
        self.ventana.configure(fg_color=t["bg_surface"])
        self.ventana.grab_set()

        self.ventana.update_idletasks()
        x = (self.ventana.winfo_screenwidth() // 2) - 425
        y = (self.ventana.winfo_screenheight() // 2) - 350
        self.ventana.geometry(f"+{x}+{y}")

        header = ctk.CTkFrame(self.ventana, fg_color="transparent")
        header.pack(fill="x", padx=20, pady=(15, 5))

        ctk.CTkLabel(header, text="📥 Importar Vocerías desde Excel",
                     font=("Arial", 14, "bold"), text_color=t["text_primary"]).pack(side="left")

        ctk.CTkButton(header, text="✕", width=30, height=30,
                      fg_color="transparent", hover_color=t["bg_hover"],
                      text_color=t["text_primary"], corner_radius=6,
                      command=self.ventana.destroy).pack(side="right")

        ctk.CTkLabel(self.ventana,
                     text="Importa vocerías en bloque subiendo un archivo .xlsx con el formato de la plantilla.\n"
                          "El archivo se agrupa por 'NOMBRE DEL COMITÉ' y cada comité debe tener entre 7 y 10 vocerías.",
                     font=("Arial", 9), text_color=t["text_muted"],
                     wraplength=750, justify="left").pack(anchor="w", padx=20, pady=(0, 10))

        self._frame_pasos = ctk.CTkFrame(self.ventana, fg_color="transparent")
        self._frame_pasos.pack(fill="both", expand=True, padx=20, pady=(0, 15))

        btn_frame = ctk.CTkFrame(self.ventana, fg_color="transparent")
        btn_frame.pack(fill="x", padx=20, pady=(0, 15))

        self._btn_cerrar = ctk.CTkButton(btn_frame, text="Cerrar",
                                         font=("Arial", 10), fg_color=t["bg_card"],
                                         hover_color=t["bg_hover"], corner_radius=8,
                                         command=self.ventana.destroy)
        self._btn_cerrar.pack(side="right")

    def _limpiar_pasos(self):
        for w in self._frame_pasos.winfo_children():
            w.destroy()

    def _mostrar_paso1(self):
        t = self.t
        self._limpiar_pasos()
        self._ruta_archivo = None
        self._resultado_previo = None

        card = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card.pack(fill="x", padx=10)

        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="x", padx=20, pady=20)

        ctk.CTkLabel(inner, text="Paso 1 de 3 — Seleccionar archivo Excel",
                     font=("Arial", 11, "bold"), text_color=t["text_primary"]).pack(anchor="w", pady=(0, 4))

        ctk.CTkLabel(inner,
                     text="El archivo debe seguir la plantilla descargable y contener: ESTADO, MUNICIPIO, PARROQUIA, ASIC, "
                          "NOMBRE DEL COMITÉ, CENTRO DISPENSADOR, NOMBRE DEL CENTRO, NOMBRES Y APELLIDO, C.I., "
                          "TELÉFONO, VOCERÍA, FECHA DE NACIMIENTO, EDAD, SEXO, PROFESIÓN/OFICIO, CORREO, FECHAS DE ELECCIÓN.",
                     font=("Arial", 8), text_color=t["text_muted"],
                     wraplength=750, justify="left").pack(anchor="w", pady=(0, 15))

        self._lbl_ruta = ctk.CTkLabel(inner, text="Ningún archivo seleccionado",
                                      font=("Arial", 9, "italic"), text_color=t["text_muted"])
        self._lbl_ruta.pack(anchor="w", pady=(0, 12))

        fila_botones = ctk.CTkFrame(inner, fg_color="transparent")
        fila_botones.pack(anchor="w")

        ctk.CTkButton(fila_botones, text="📂 Examinar…",
                      font=("Arial", 9), fg_color=t["bg_card"],
                      hover_color=t["bg_hover"], text_color=t["text_primary"],
                      corner_radius=8, command=self._seleccionar_archivo).pack(side="left", padx=(0, 10))

        self._btn_previsualizar = ctk.CTkButton(
            fila_botones, text="▶ Previsualizar",
            font=("Arial", 9, "bold"), fg_color=t["bg_hover"],
            hover_color=t["bg_hover"], text_color=t["text_dim"],
            text_color_disabled=t["text_dim"], corner_radius=8,
            state="disabled", command=self._iniciar_previo)
        self._btn_previsualizar.pack(side="left")

    def _seleccionar_archivo(self):
        t = self.t
        ruta = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=[("Archivos Excel", "*.xlsx *.xls"), ("Todos los archivos", "*.*")],
            parent=self.ventana
        )
        if ruta:
            self._ruta_archivo = ruta
            nombre = os.path.basename(ruta)
            self._lbl_ruta.configure(text=f"📄 {nombre}  ({ruta})", text_color=t["text_primary"])
            self._btn_previsualizar.configure(
                state="normal", fg_color=t["accent"],
                hover_color=t["accent_dim"], text_color=t["text_primary"])

    def _iniciar_previo(self):
        self._btn_previsualizar.configure(state="disabled", text="Procesando…")
        self.ventana.configure(cursor="wait")

        def _tarea():
            try:
                svc = ExcelServiceVoceria(id_usuario_registro=self.usuario["id_usuario"])
                resultado = svc.previsualizar(self._ruta_archivo)
                self.ventana.after(0, lambda: self._mostrar_paso2(resultado))
            except Exception as exc:
                self.ventana.after(0, lambda e=exc: self._error_previo(e))

        threading.Thread(target=_tarea, daemon=True).start()

    def _error_previo(self, exc):
        self.ventana.configure(cursor="")
        messagebox.showerror("Error al leer el archivo",
                             f"No se pudo procesar el archivo:\n\n{exc}", parent=self.ventana)
        self._btn_previsualizar.configure(state="normal", text="▶ Previsualizar",
                                          fg_color=self.t["bg_hover"],
                                          text_color=self.t["text_dim"])

    def _mostrar_paso2(self, resultado):
        t = self.t
        self.ventana.configure(cursor="")
        self._resultado_previo = resultado
        self._limpiar_pasos()

        comites_validos = resultado["comites_validos"]
        comites_error = resultado["comites_error"]
        filas_error_sin_comite = resultado.get("filas_error_sin_comite", [])
        total_filas = resultado["total_filas_procesadas"]

        card_resumen = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card_resumen.pack(fill="x", pady=(0, 10), padx=10)

        inner = ctk.CTkFrame(card_resumen, fg_color="transparent")
        inner.pack(fill="x", padx=15, pady=12)

        ctk.CTkLabel(inner, text="Paso 2 de 3 — Resultado del análisis",
                     font=("Arial", 11, "bold"), text_color=t["text_primary"]).pack(anchor="w", pady=(0, 8))

        fila_resumen = ctk.CTkFrame(inner, fg_color="transparent")
        fila_resumen.pack(anchor="w")

        def _chip(parent, texto, color):
            chip = ctk.CTkFrame(parent, fg_color=color, corner_radius=6)
            chip.pack(side="left", padx=(0, 8))
            ctk.CTkLabel(chip, text=texto, font=("Arial", 8, "bold"),
                         text_color=t["text_primary"]).pack(padx=8, pady=3)

        _chip(fila_resumen, f"📄 Total filas: {total_filas}", t["bg_card"])
        _chip(fila_resumen, f"✅ Comités válidos: {len(comites_validos)}", t["success"])
        _chip(fila_resumen, f"❌ Comités con errores: {resultado['total_comites_error']}",
              t["red"] if comites_error else t["bg_hover"])
        if filas_error_sin_comite:
            _chip(fila_resumen, f"⚠ Filas sin comité: {len(filas_error_sin_comite)}", t["amber"])

        if comites_validos:
            self._seccion_comites_validos(comites_validos)
        if comites_error:
            self._seccion_comites_error(comites_error)
        if filas_error_sin_comite:
            self._seccion_filas_sin_comite(filas_error_sin_comite)

        card_acciones = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card_acciones.pack(fill="x", padx=10, pady=(8, 0))

        inner_acc = ctk.CTkFrame(card_acciones, fg_color="transparent")
        inner_acc.pack(fill="x", padx=15, pady=12)

        ctk.CTkButton(inner_acc, text="← Seleccionar otro archivo",
                      font=("Arial", 9), fg_color=t["bg_card"],
                      hover_color=t["bg_hover"], text_color=t["text_primary"],
                      corner_radius=8, command=self._mostrar_paso1).pack(side="left", padx=(0, 10))

        if comites_validos:
            total_vocerias = sum(len(c["filas"]) for c in comites_validos)
            ctk.CTkButton(inner_acc,
                          text=f"✔ Importar {len(comites_validos)} comités ({total_vocerias} vocerías)",
                          font=("Arial", 9, "bold"), fg_color=t["success"],
                          hover_color=t["success_dim"], text_color=t["text_primary"],
                          corner_radius=8, command=self._confirmar_importar).pack(side="left")
        else:
            ctk.CTkLabel(inner_acc,
                         text="No hay comités válidos para importar. Corrige el archivo y vuelve a intentarlo.",
                         font=("Arial", 8, "italic"), text_color=t["red"]).pack(side="left")

    def _seccion_comites_validos(self, comites):
        t = self.t
        marco = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        marco.pack(fill="both", expand=True, pady=(0, 8), padx=10)

        inner = ctk.CTkFrame(marco, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=10, pady=8)

        ctk.CTkLabel(inner, text="✅ Comités listos para importar", font=("Arial", 9, "bold"),
                     text_color=t["success"]).pack(anchor="w", pady=(0, 4))

        tree_frame = ctk.CTkFrame(inner, fg_color="transparent")
        tree_frame.pack(fill="both", expand=True)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("ImportVoceria.Treeview",
                        background=t["bg_base"], foreground=t["text_primary"],
                        fieldbackground=t["bg_base"], rowheight=24)
        style.configure("ImportVoceria.Treeview.Heading",
                        background=t["bg_card"], foreground=t["text_primary"],
                        font=("Arial", 9, "bold"), relief="flat")
        style.map("ImportVoceria.Treeview",
                  background=[("selected", t["accent"])],
                  foreground=[("selected", t["text_primary"])])

        tree = ttk.Treeview(tree_frame, columns=("comite", "vocerias", "asic", "municipio"), show="headings",
                            style="ImportVoceria.Treeview", height=min(len(comites), 8))

        tree.heading("comite", text="Nombre del Comité")
        tree.heading("vocerias", text="Vocerías")
        tree.heading("asic", text="ASIC")
        tree.heading("municipio", text="Municipio")
        tree.column("comite", width=300, minwidth=200, anchor="w")
        tree.column("vocerias", width=80, minwidth=60, anchor="center")
        tree.column("asic", width=180, minwidth=120, anchor="w")
        tree.column("municipio", width=180, minwidth=120, anchor="w")

        for comite in comites:
            primera = comite["filas"][0]["datos"]
            tree.insert("", "end", values=(
                comite["nombre_comite"],
                len(comite["filas"]),
                primera.get("asic", ""),
                primera.get("municipio", ""),
            ))

        scrolly = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        scrollx = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=scrolly.set, xscrollcommand=scrollx.set)
        scrolly.pack(side="right", fill="y")
        scrollx.pack(side="bottom", fill="x")
        tree.pack(fill="both", expand=True)

    def _seccion_comites_error(self, comites_error):
        t = self.t
        marco = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        marco.pack(fill="x", pady=(0, 8), padx=10)

        inner = ctk.CTkFrame(marco, fg_color="transparent")
        inner.pack(fill="x", padx=10, pady=8)

        ctk.CTkLabel(inner, text="❌ Comités con errores (no se importarán)", font=("Arial", 9, "bold"),
                     text_color=t["red"]).pack(anchor="w", pady=(0, 4))

        tree_frame = ctk.CTkFrame(inner, fg_color="transparent")
        tree_frame.pack(fill="x")

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("ImportVoceriaError.Treeview",
                        background=t["bg_base"], foreground=t["text_primary"],
                        fieldbackground=t["bg_base"], rowheight=24)
        style.configure("ImportVoceriaError.Treeview.Heading",
                        background=t["bg_card"], foreground=t["text_primary"],
                        font=("Arial", 9, "bold"), relief="flat")
        style.map("ImportVoceriaError.Treeview",
                  background=[("selected", t["accent"])],
                  foreground=[("selected", t["text_primary"])])

        comites_agrupados = {}
        for err in comites_error:
            nombre = err["nombre_comite"]
            if nombre not in comites_agrupados:
                comites_agrupados[nombre] = []
            comites_agrupados[nombre].extend(err["errores"])

        tree = ttk.Treeview(tree_frame, columns=("comite", "errores"), show="headings",
                            style="ImportVoceriaError.Treeview", height=min(len(comites_agrupados), 6))

        tree.heading("comite", text="Nombre del Comité")
        tree.heading("errores", text="Errores detectados")
        tree.column("comite", width=300, minwidth=200, anchor="w")
        tree.column("errores", width=450, minwidth=300, anchor="w")

        for nombre, errores in comites_agrupados.items():
            errores_unicos = list(dict.fromkeys(errores))
            tree.insert("", "end", values=(nombre, " | ".join(errores_unicos)))

        scrolly = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        scrollx = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=scrolly.set, xscrollcommand=scrollx.set)
        scrolly.pack(side="right", fill="y")
        scrollx.pack(side="bottom", fill="x")
        tree.pack(fill="x")

    def _seccion_filas_sin_comite(self, filas_error):
        t = self.t
        marco = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        marco.pack(fill="x", pady=(0, 8), padx=10)

        inner = ctk.CTkFrame(marco, fg_color="transparent")
        inner.pack(fill="x", padx=10, pady=8)

        ctk.CTkLabel(inner, text="⚠ Filas sin comité válido (no se importarán)", font=("Arial", 9, "bold"),
                     text_color=t["amber"]).pack(anchor="w", pady=(0, 4))

        tree_frame = ctk.CTkFrame(inner, fg_color="transparent")
        tree_frame.pack(fill="x")

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("ImportVoceriaError.Treeview",
                        background=t["bg_base"], foreground=t["text_primary"],
                        fieldbackground=t["bg_base"], rowheight=24)
        style.configure("ImportVoceriaError.Treeview.Heading",
                        background=t["bg_card"], foreground=t["text_primary"],
                        font=("Arial", 9, "bold"), relief="flat")

        tree = ttk.Treeview(tree_frame, columns=("fila", "comite", "errores"), show="headings",
                            style="ImportVoceriaError.Treeview", height=min(len(filas_error), 6))

        tree.heading("fila", text="# Fila")
        tree.heading("comite", text="Comité")
        tree.heading("errores", text="Errores")
        tree.column("fila", width=60, minwidth=50, anchor="center")
        tree.column("comite", width=200, minwidth=150, anchor="w")
        tree.column("errores", width=500, minwidth=300, anchor="w")

        for err in filas_error:
            tree.insert("", "end", values=(
                err["numero_fila"],
                err["nombre_comite"],
                " | ".join(err["errores"])
            ))

        scrolly = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        scrollx = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=scrolly.set, xscrollcommand=scrollx.set)
        scrolly.pack(side="right", fill="y")
        scrollx.pack(side="bottom", fill="x")
        tree.pack(fill="x")

    def _confirmar_importar(self):
        comites_validos = self._resultado_previo["comites_validos"]
        total_vocerias = sum(len(c["filas"]) for c in comites_validos)
        if messagebox.askyesno("Confirmar importación",
                               f"Se importarán {len(comites_validos)} comités ({total_vocerias} vocerías) en la base de datos.\n\n"
                               f"¿Deseas continuar?",
                               parent=self.ventana):
            self._ejecutar_importacion(comites_validos)

    def _ejecutar_importacion(self, comites_validos):
        self.ventana.configure(cursor="wait")

        def _tarea():
            try:
                svc = ExcelServiceVoceria(id_usuario_registro=self.usuario["id_usuario"])
                resultado = svc.importar(comites_validos)
                self.ventana.after(0, lambda r=resultado: self._mostrar_paso3(r))
            except Exception as exc:
                self.ventana.after(0, lambda e=exc: self._error_importar(e))

        threading.Thread(target=_tarea, daemon=True).start()

    def _error_importar(self, exc):
        self.ventana.configure(cursor="")
        messagebox.showerror("Error al importar",
                             f"Ocurrió un error durante la importación:\n\n{exc}", parent=self.ventana)

    def _mostrar_paso3(self, resultado):
        t = self.t
        self.ventana.configure(cursor="")
        self._limpiar_pasos()

        comites_insertados = resultado["comites_insertados"]
        voceros_insertados = resultado["voceros_insertados"]
        errores_bd = resultado["errores_al_insertar"]

        card = ctk.CTkFrame(self._frame_pasos, fg_color=t["bg_surface"], corner_radius=10)
        card.pack(fill="x", padx=10)

        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="x", padx=25, pady=25)

        icono = "✅" if not errores_bd else "⚠️"
        color = t["success"] if not errores_bd else t["amber"]
        titulo = "Importación completada" if not errores_bd else "Importación completada con advertencias"

        ctk.CTkLabel(inner, text=f"{icono}  {titulo}",
                     font=("Arial", 13, "bold"), text_color=color).pack(anchor="w", pady=(0, 10))

        ctk.CTkLabel(inner, text=f"Se insertaron correctamente {comites_insertados} comités ({voceros_insertados} vocerías).",
                     font=("Arial", 10), text_color=t["text_primary"]).pack(anchor="w")

        if errores_bd:
            ctk.CTkLabel(inner,
                         text=f"{len(errores_bd)} comités no pudieron insertarse (ver detalles abajo).",
                         font=("Arial", 9), text_color=t["red"]).pack(anchor="w", pady=(4, 0))
            for err in errores_bd:
                ctk.CTkLabel(inner, text=f"  • {err['nombre_comite']}: {err['error']}",
                             font=("Arial", 8), text_color=t["red"]).pack(anchor="w")

        ctk.CTkButton(inner, text="↩ Hacer otra importación",
                      font=("Arial", 9), fg_color=t["bg_card"],
                      hover_color=t["bg_hover"], text_color=t["text_primary"],
                      corner_radius=8, command=self._mostrar_paso1).pack(anchor="w", pady=(18, 0))

        ctk.CTkButton(inner, text="Cerrar",
                      font=("Arial", 9, "bold"), fg_color=t["accent"],
                      hover_color=t["accent_dim"], text_color=t["text_primary"],
                      corner_radius=8, command=lambda: (self.ventana.destroy(), self.parent_view.controlador.cargar_datos_tabla())
                      ).pack(anchor="w", pady=(8, 0))


class ControladorVoceria:
    def __init__(self, parent):
        self.modelo = ModeloVoceria()

        self.asic_map = {}
        self.tipo_map = {}
        self.tipo_voceria_map = {}
        self.tipo_voceria_map_normalizada = {}
        self._cargar_catalogos()

        self.vista = VistaVoceria(parent, self)
        self.cargar_datos_tabla()

    @staticmethod
    def _normalizar_texto(valor):
        if valor is None:
            return ""
        valor = unicodedata.normalize("NFKD", str(valor))
        valor = valor.encode("ascii", "ignore").decode("ascii")
        valor = valor.lower().strip()
        valor = re.sub(r"\s+", " ", valor)
        return valor

    def _cargar_catalogos(self):
        try:
            asics = self.modelo.listar_asics()
            self.asic_map = {a["nombre"]: a["id_asic"] for a in asics}

            tipos = self.modelo.listar_tipos_voceria()
            self.tipo_map = {t["nombre"]: t["id_tipo"] for t in tipos}
            self.tipo_voceria_map = {t["nombre"]: t["id_tipo"] for t in tipos}
            self.tipo_voceria_map_normalizada = {
                self._normalizar_texto(nombre): id_tipo
                for nombre, id_tipo in self.tipo_voceria_map.items()
            }
        except Exception as e:
            messagebox.showerror("Error de base de datos", f"No se pudieron cargar los catálogos: {str(e)}")

    def aplicar_filtros(self):
        asic_sel = self.vista.combo_filtro_asic.get()
        tipo_sel = self.vista.combo_filtro_tipo.get()
        f_desde = self.vista.ent_filtro_desde.get().strip()
        f_hasta = self.vista.ent_filtro_hasta.get().strip()

        filtros = {}

        if asic_sel != "Todos":
            filtros["id_asic"] = self.asic_map.get(asic_sel)
        if tipo_sel != "Todos":
            filtros["id_tipo"] = self.tipo_map.get(tipo_sel)

        if f_desde and f_desde != "AAAA-MM-DD":
            try:
                datetime.strptime(f_desde, "%Y-%m-%d")
                filtros["fecha_desde"] = f_desde
            except ValueError:
                messagebox.showerror("Error en Filtros", "La fecha 'Desde' debe tener el formato AAAA-MM-DD (ej: 2026-05-17).")
                return

        if f_hasta and f_hasta != "AAAA-MM-DD":
            try:
                datetime.strptime(f_hasta, "%Y-%m-%d")
                filtros["fecha_hasta"] = f_hasta
            except ValueError:
                messagebox.showerror("Error en Filtros", "La fecha 'Hasta' debe tener el formato AAAA-MM-DD (ej: 2026-05-17).")
                return

        self.cargar_datos_tabla(filtros)

    def limpiar_filtros(self):
        self.vista.combo_filtro_asic.set("Todos")
        self.vista.combo_filtro_tipo.set("Todos")
        self.vista.ent_filtro_desde.set_date("")
        self.vista.ent_filtro_hasta.set_date("")
        self.cargar_datos_tabla()

    def cargar_datos_tabla(self, filtros=None):
        for item in self.vista.tabla.get_children():
            self.vista.tabla.delete(item)

        try:
            lista = self.modelo.listar_vocerias(filtros)
            for voc in lista:
                f_ini = voc.get("fecha_registro") or voc.get("fecha")
                if isinstance(f_ini, (date, datetime)):
                    f_ini = f_ini.strftime("%d/%m/%Y")

                cantidad = voc.get("cantidad_vocerias", 0)

                self.vista.tabla.insert("", "end", iid=str(voc["id_voceria"]), values=(
                    voc["id_voceria"],
                    voc.get("asic_nombre", ""),
                    f_ini,
                    cantidad,
                    voc.get("estado", ""),
                    voc.get("municipio", ""),
                    voc.get("parroquia", ""),
                ))
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al cargar las vocerías: {str(e)}")

    def abrir_form_nueva(self):
        self.vista.abrir_formulario()

    def editar_seleccionada(self):
        seleccion = self.vista.tabla.selection()
        if not seleccion:
            messagebox.showwarning("Atención", "Selecciona una vocería de la tabla para editar.")
            return
        id_voceria = int(seleccion[0])
        try:
            voc = self.modelo.obtener_voceria_completa(id_voceria)
            if not voc:
                messagebox.showerror("Error", "No se pudo recuperar la información de la vocería.")
                return
            self.vista.abrir_formulario(voc)
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al obtener la vocería: {str(e)}")

    def eliminar_seleccionada(self):
        seleccion = self.vista.tabla.selection()
        if not seleccion:
            messagebox.showwarning("Atención", "Selecciona una vocería de la tabla para eliminar.")
            return
        id_voceria = int(seleccion[0])
        valores = self.vista.tabla.item(seleccion[0])["values"]
        tipo_act = valores[1]
        fecha_act = valores[2]

        if messagebox.askyesno("Confirmar eliminación",
                               f"¿Deseas eliminar la vocería '{tipo_act}' del día {fecha_act}?\nEsta acción no se puede deshacer."):
            try:
                ok, msg = self.modelo.eliminar_voceria(id_voceria)
                if ok:
                    messagebox.showinfo("Éxito", msg)
                    self.cargar_datos_tabla()
                else:
                    messagebox.showerror("Error", msg)
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo eliminar la vocería: {str(e)}")

    def on_seleccion_cambio(self, event=None):
        seleccion = self.vista.tabla.selection()
        estado = "normal" if seleccion else "disabled"
        self.vista.btn_editar.configure(state=estado)
        self.vista.btn_eliminar.configure(state=estado)

    def abrir_detalle_voceria(self, event=None):
        seleccion = self.vista.tabla.selection()
        if not seleccion:
            return
        id_voceria = int(seleccion[0])

        try:
            voc = self.modelo.obtener_voceria_completa(id_voceria)
            if not voc:
                messagebox.showerror("Error", "No se pudo recuperar la información de la vocería.")
                return
            self.vista.mostrar_detalle_voceria(voc)
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al obtener el detalle: {str(e)}")

    def guardar_voceria(self, ventana, lbl_error, datos_edicion,
                        estado, municipio, parroquia, asic,
                        centro_dispensador, nombre_centro, nombre_comite,
                        items_validos):
        id_asic = self.asic_map.get(asic)

        items = []
        for item in items_validos:
            w = item["widgets"]
            valor_voceria = (w["voceria"].get() or "").strip()
            id_tipo_voceria = (
                self.tipo_voceria_map.get(valor_voceria)
                or self.tipo_voceria_map_normalizada.get(self._normalizar_texto(valor_voceria))
            )

            item_data = {
                "nombres": w["nombres"].get().strip(),
                "ci": w["ci"].get().strip(),
                "telefono": w["tel"].get().strip(),
                "fecha_nacimiento": w["fnac"].get().strip(),
                "edad": w["edad"].get().strip(),
                "id_tipo_voceria": id_tipo_voceria,
                "sexo": w["sexo"].get(),
                "profesion": w["prof"].get().strip(),
                "correo": w["correo"].get().strip(),
                "fechas_eleccion": w["felecc"].get().strip(),
            }
            items.append(item_data)

        id_tipo = None
        if items:
            id_tipo = items[0].get("id_tipo_voceria") or next(iter(self.tipo_voceria_map.values()), None)

        id_voceria_guardar = datos_edicion.get("id_voceria") if datos_edicion else None

        ok, msg = self.modelo.guardar_voceria(
            id_asic=id_asic,
            id_tipo=id_tipo,
            estado=estado,
            municipio=municipio,
            parroquia=parroquia,
            centro_dispensador=centro_dispensador,
            nombre_centro=nombre_centro,
            nombre_comite=nombre_comite,
            items=items,
            id_voceria=id_voceria_guardar,
        )
        if ok:
            messagebox.showinfo("Éxito", msg)
            ventana.destroy()
            self.cargar_datos_tabla()
        else:
            lbl_error.configure(text=msg)

    def descargar_plantilla(self):
        if not os.path.exists(VOCERIA_TEMPLATE_PATH):
            try:
                build_voceria_template(VOCERIA_TEMPLATE_PATH)
            except Exception as e:
                messagebox.showerror("Plantilla no disponible",
                                     f"No se pudo crear la plantilla de vocerías:\n{str(e)}")
                return

        destino = filedialog.asksaveasfilename(
            title="Guardar plantilla de vocerías como...",
            defaultextension=".xlsx",
            initialfile="plantilla_importacion_voceria.xlsx",
            filetypes=[("Archivos Excel", "*.xlsx")],
        )
        if destino:
            shutil.copy2(VOCERIA_TEMPLATE_PATH, destino)
            messagebox.showinfo("Plantilla guardada",
                                f"Plantilla de vocerías guardada en:\n{destino}")

    def abrir_modal_importar(self):
        _ImportarVoceriaModal(self.vista)

    def exportar_consolidado(self):
        from tkinter import filedialog
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        import os

        try:
            filtros = {}
            asic_sel = self.vista.combo_filtro_asic.get()
            tipo_sel = self.vista.combo_filtro_tipo.get()
            f_desde = self.vista.ent_filtro_desde.get().strip()
            f_hasta = self.vista.ent_filtro_hasta.get().strip()

            if asic_sel != "Todos":
                filtros["id_asic"] = self.asic_map.get(asic_sel)
            if tipo_sel != "Todos":
                filtros["id_tipo"] = self.tipo_map.get(tipo_sel)
            if f_desde and f_desde != "AAAA-MM-DD":
                filtros["fecha_desde"] = f_desde
            if f_hasta and f_hasta != "AAAA-MM-DD":
                filtros["fecha_hasta"] = f_hasta

            lista_comites = self.modelo.listar_vocerias(filtros)

            if not lista_comites:
                messagebox.showinfo("Sin datos", "No hay vocerías para exportar con los filtros actuales.")
                return

            destino = filedialog.asksaveasfilename(
                title="Guardar consolidado de vocerías como...",
                defaultextension=".xlsx",
                initialfile="Consolidado_Vocerias_CS.xlsx",
                filetypes=[("Archivos Excel", "*.xlsx")],
            )
            if not destino:
                return

            template_path = os.path.join(os.path.dirname(__file__), "..", "..", "assets", "templates", "plantilla_importacion_voceria.xlsx")
            template_path = os.path.normpath(template_path)

            if not os.path.exists(template_path):
                messagebox.showerror("Error", f"No se encuentra la plantilla en:\n{template_path}")
                return

            wb = openpyxl.load_workbook(template_path)
            ws = wb.active

            row_num = 3
            counter = 1

            for comite_idx, comite in enumerate(lista_comites):
                id_comite = comite.get("id_voceria")
                if not id_comite:
                    continue

                try:
                    comite_completo = self.modelo.obtener_voceria_completa(id_comite)
                    if not comite_completo:
                        continue

                    items = comite_completo.get("items", [])
                    if not items:
                        continue

                    estado = comite_completo.get("estado", "")
                    municipio = comite_completo.get("municipio", "")
                    parroquia = comite_completo.get("parroquia", "")
                    asic_nombre = comite_completo.get("asic_nombre", "")
                    nombre_comite = comite_completo.get("nombre_comite", "")
                    centro_dispensador = comite_completo.get("centro_dispensador", "")
                    nombre_centro = comite_completo.get("nombre_centro", "")

                    for item in items:
                        fnac = item.get("fecha_nacimiento")
                        if isinstance(fnac, (date, datetime)):
                            fnac = fnac.strftime("%d/%m/%Y")

                        felecc = item.get("fechas_eleccion")
                        if isinstance(felecc, (date, datetime)):
                            felecc = felecc.strftime("%d/%m/%Y")

                        ws.cell(row=row_num, column=1, value=counter)
                        ws.cell(row=row_num, column=2, value=estado)
                        ws.cell(row=row_num, column=3, value=municipio)
                        ws.cell(row=row_num, column=4, value=parroquia)
                        ws.cell(row=row_num, column=5, value=asic_nombre)
                        ws.cell(row=row_num, column=6, value=nombre_comite)
                        ws.cell(row=row_num, column=7, value=centro_dispensador)
                        ws.cell(row=row_num, column=8, value=nombre_centro)
                        ws.cell(row=row_num, column=9, value=item.get("nombres", ""))
                        ws.cell(row=row_num, column=10, value=item.get("ci", ""))
                        ws.cell(row=row_num, column=11, value=item.get("telefono", ""))
                        ws.cell(row=row_num, column=12, value=item.get("tipo_voceria_nombre", ""))
                        ws.cell(row=row_num, column=13, value=fnac or "")
                        ws.cell(row=row_num, column=14, value=item.get("edad", ""))
                        ws.cell(row=row_num, column=15, value=item.get("sexo", ""))
                        ws.cell(row=row_num, column=16, value=item.get("profesion", ""))
                        ws.cell(row=row_num, column=17, value=item.get("correo", ""))
                        ws.cell(row=row_num, column=18, value=felecc or "")

                        row_num += 1
                        counter += 1

                    # Add blank row between committees (except after the last one)
                    if comite_idx < len(lista_comites) - 1:
                        row_num += 1

                except Exception as e:
                    print(f"Error procesando comité {id_comite}: {e}")
                    continue

            wb.save(destino)
            messagebox.showinfo("Éxito", f"Consolidado exportado correctamente a:\n{destino}")

        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al exportar: {str(e)}")

    def _obtener_vocerias_usadas_por_otros(self, item_actual, items_voceria):
        usadas = set()
        for item in items_voceria:
            if item != item_actual and item["frame"].winfo_exists():
                val = item["widgets"]["voceria"].get()
                if val and val not in ("(Sin opciones disponibles)", "Seleccionar vocería..."):
                    usadas.add(val)
        return usadas

    def _actualizar_dropdowns_voceria(self, items_voceria, vocerias_disponibles):
        for item in items_voceria:
            if not item["frame"].winfo_exists():
                continue
            widget = item["widgets"]["voceria"]
            val_actual = widget.get()
            otras_usadas = self._obtener_vocerias_usadas_por_otros(item, items_voceria)
            disponibles_para_este_item = [v for v in vocerias_disponibles if v not in otras_usadas]

            valores = ["Seleccionar vocería..."] + disponibles_para_este_item
            widget.configure(values=valores, state="normal")
            if val_actual in valores:
                widget.set(val_actual)
            else:
                widget.set("Seleccionar vocería...")

    def _actualizar_estado_botones_formulario(self):
        if not hasattr(self.vista, "items_voceria"):
            return

        items_activos = [i for i in self.vista.items_voceria if i["frame"].winfo_exists()]
        count = len(items_activos)

        generales_llenos = self._verificar_generales_llenos()
        todos_items_llenos = items_activos and all(self._verificar_item_completo(i) for i in items_activos)

        self._actualizar_bordes_items(items_activos)

        for item in items_activos:
            item["btn_eliminar"].configure(state="normal")

        if hasattr(self.vista, 'btn_agregar_item') and self.vista.btn_agregar_item.winfo_exists():
            self.vista.btn_agregar_item.configure(state="normal" if count < 10 else "disabled")

        if hasattr(self.vista, 'btn_guardar') and self.vista.btn_guardar.winfo_exists():
            self.vista.btn_guardar.configure(state="normal")
            if count >= 7 and count <= 10 and generales_llenos and todos_items_llenos:
                self.vista.lbl_nota_min.configure(text="✓ Todos los campos obligatorios están completos", text_color=self.vista.t["success"])
            elif count < 7:
                self.vista.lbl_nota_min.configure(text=f"⚠ Debe agregar al menos 7 vocerías (actual: {count})", text_color=self.vista.t["red"])
            elif count > 10:
                self.vista.lbl_nota_min.configure(text=f"⚠ Máximo 10 vocerías permitidas (actual: {count})", text_color=self.vista.t["red"])
            elif count >= 1 and not generales_llenos:
                self.vista.lbl_nota_min.configure(text="⚠ Complete los datos generales obligatorios (*)", text_color=self.vista.t["amber"])
            elif count >= 1 and not todos_items_llenos:
                self.vista.lbl_nota_min.configure(text="⚠ Complete todos los campos de cada vocería (*)", text_color=self.vista.t["amber"])
            else:
                self.vista.lbl_nota_min.configure(text="⚠ Debe agregar al menos 7 vocerías (máx. 10)", text_color=self.vista.t["text_muted"])

    def _verificar_generales_llenos(self):
        estado = self.vista.ent_estado.get().strip() if hasattr(self.vista, "ent_estado") else ""
        municipio = self.vista.ent_municipio.get().strip() if hasattr(self.vista, "ent_municipio") else ""
        parroquia = self.vista.ent_parroquia.get().strip() if hasattr(self.vista, "ent_parroquia") else ""
        asic = self.vista.ent_asic.get().strip() if hasattr(self.vista, "ent_asic") else ""
        centro_dispensador = self.vista.ent_centro_dispensador.get().strip() if hasattr(self.vista, "ent_centro_dispensador") else ""
        nombre_centro = self.vista.ent_nombre_centro.get().strip() if hasattr(self.vista, "ent_nombre_centro") else ""
        nombre_comite = self.vista.ent_nombre_comite.get().strip() if hasattr(self.vista, "ent_nombre_comite") else ""

        if not estado or estado in ("Seleccionar...", ""):
            return False
        if not municipio or municipio in ("Seleccionar...", ""):
            return False
        if not parroquia or parroquia == "Parroquia":
            return False
        if not asic or asic not in self.asic_map:
            return False
        if not centro_dispensador:
            return False
        if not nombre_centro:
            return False
        if not nombre_comite:
            return False
        return True

    def _verificar_item_completo(self, item):
        if not item["frame"].winfo_exists():
            return False
        w = item["widgets"]
        campos_obligatorios = [
            "nombres", "ci", "tel", "fnac", "edad", "voceria", "sexo", "prof", "correo", "felecc"
        ]
        for key in campos_obligatorios:
            widget = w.get(key)
            if not widget:
                return False
            val = widget.get().strip() if hasattr(widget, "get") else ""
            if not val or val in ("(Sin opciones disponibles)", "Seleccionar vocería...", "AAAA-MM-DD", "Edad"):
                return False
            if key == "correo" and "@" not in val:
                return False
            if key == "edad":
                if not val.isdigit() or len(val) > 3 or int(val) > 999:
                    return False
        return True

    def _actualizar_bordes_items(self, items_activos):
        for item in items_activos:
            if not item["frame"].winfo_exists():
                continue
            completo = self._verificar_item_completo(item)
            color = "#317E87" if completo else self.vista.t["border"]
            item["frame"].configure(border_color=color)

    def _renumerar_items(self, items_voceria):
        for idx, item in enumerate(items_voceria):
            if item["frame"].winfo_exists():
                item["index"] = idx + 1
                item["lbl_titulo"].configure(text=f"Vocería #{idx + 1}")
                item["btn_eliminar"].configure(state="normal")
                self.vista._actualizar_header_item(item)

    def agregar_siguiente_voceria(self, frame_items, items_voceria, vocerias_disponibles):
        self.vista._crear_item_voceria(frame_items, items_voceria, vocerias_disponibles, expandido=True)
        self._actualizar_dropdowns_voceria(items_voceria, vocerias_disponibles)
        self._actualizar_estado_botones_formulario()
        self._renumerar_items(items_voceria)

    def _eliminar_item_voceria(self, item_data, items_voceria, vocerias_disponibles):
        if item_data in items_voceria:
            item_data["frame"].destroy()
            items_voceria.remove(item_data)
            self._actualizar_dropdowns_voceria(items_voceria, vocerias_disponibles)
            self._actualizar_estado_botones_formulario()
            self._renumerar_items(items_voceria)


__all__ = ["ControladorVoceria"]
