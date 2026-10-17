import os
import openpyxl
import re
import unicodedata
from datetime import date, datetime, timedelta

from constants import (
    ACTIVIDAD_ESTABLECIMIENTO_OPCIONES,
    MUNICIPIOS_TACHIRA,
    UBICACIONES_ASIC_PREDETERMINADAS,
    VOCERIA_CENTRO_DISPENSADOR_OPCIONES,
    VOCERIA_ESTADO_OPCIONES,
    VOCERIA_SEXO_OPCIONES,
    VOCERIA_TEMPLATE_PATH,
    VOCERIA_TIPO_OPCIONES,
)
from ...daos import (
    ActividadDAO,
    AsicDAO,
    MunicipioDAO,
    ParroquiaDAO,
    TipoActividadDAO,
    TipoVoceriaDAO,
)
from ...validators import validar_archivo_excel, validar_correo

COLUMNAS_REQUERIDAS = [
    "asic", "tipo_actividad", "fecha", "tema",
    "participantes",
]

HEADER_ALIASES = {
    "fecha": "fecha",
    "fecha inicio": "fecha",
    "fecha inicio yyyy mm dd": "fecha",
    "fecha fin": "fecha",
    "estado": "estado",
    "municipio": "municipio",
    "parroquia": "parroquia",
    "asic": "asic",
    "establecimiento de salud": "establecimiento_salud",
    "establecimiento de salud cdi sri hospital cat clinica popular": "establecimiento_salud",
    "nombre del establecimiento de salud": "nombre_establecimiento_salud",
    "comites de salud": "comites_de_salud",
    "comités de salud": "comites_de_salud",
    "cantidad de comites de salud": "comites_de_salud",
    "cantidad de comités de salud": "comites_de_salud",
    "cantidad de comites de salud 0 99": "comites_de_salud",
    "ubicación": "ubicacion",
    "ubicacion": "ubicacion",
    "actividad desarrollada": "tipo_actividad",
    "actividad": "tipo_actividad",
    "tipo de actividad": "tipo_actividad",
    "tipo actividad": "tipo_actividad",
    "tipo de actividad nombre exacto": "tipo_actividad",
    "asic nombre exacto": "asic",
    "tema tratado": "tema",
    "tema": "tema",
    "participantes": "participantes",
    "participantes numero": "participantes",
    "medicos": "medicos",
    "promotores en salud": "promotores_en_salud",
    "voceros de comité de salud": "voceros_de_comite_de_salud",
    "voceros de comite de salud": "voceros_de_comite_de_salud",
    "vocero del consejo de direccion del asic": "vocero_del_consejo_de_direccion_del_asic",
    "coordinador de docencia e investigacion del asic": "coordinador_de_docencia_e_investigacion_del_asic",
    "autoridades presentes": "autoridades_presentes",
    "otros actores involucrados": "otros_actores_involucrados",
    "cantidad de personas de la comunidad": "cantidad_personas_comunidad",
    "cantidad de personas de la comunidad estudiantes ninos y adultos": "cantidad_personas_comunidad",
    "estudiantes aspirantes para la salud involucrados en la actividad": "estudiantes_aspirantes",
    "total de asistencia": "total_asistencia",
}

PARTICIPANTES_CATEGORIAS = [
    "medicos", "promotores_en_salud", "voceros_de_comite_de_salud",
    "vocero_del_consejo_de_direccion_del_asic", "coordinador_de_docencia_e_investigacion_del_asic",
    "autoridades_presentes", "otros_actores_involucrados", "cantidad_personas_comunidad",
    "estudiantes_aspirantes",
]

def build_voceria_template(output_path=None, ubicaciones_asic=None):
    """Genera la plantilla Excel para vocerías CS con listas desplegables."""
    if output_path is None:
        output_path = VOCERIA_TEMPLATE_PATH

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    ws = wb.active
    ws.title = "hoja1"

    RED_FILL = "C00000"
    WHITE_FONT = "FFFFFFFF"
    thin_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    title_fill = PatternFill("solid", fgColor=RED_FILL)
    title_font = Font(bold=True, color=WHITE_FONT, size=12, name="Calibri")
    title_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_fill = PatternFill("solid", fgColor=RED_FILL)
    header_font = Font(bold=True, color=WHITE_FONT, size=12, name="Calibri")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A1:R1")
    ws["A1"] = "DATAS DE LOS VOCEROS Y VOCERAS DE LOS COMITES DE SALUD  2026"
    ws["A1"].fill = title_fill
    ws["A1"].font = title_font
    ws["A1"].alignment = title_align
    ws.row_dimensions[1].height = 36.75

    headers = [
        "Nro", "ESTADO", "MUNICIPIO", "PARROQUIA", "ASIC",
        "NOMBRE DEL COMITÉ",
        "INDIQUE QUE CENTRO DISPENSADOR DE SALUD ES  ( SRI, HOSPITAL, CAT, CDI, CLINICA POPULAR, CONSULTORIO MEDICO POPULAR, AMBULATORIO U OTRO )",
        "NOMBRE DEL CENTRO DISPENSADOR DE SALUD",
        "NOMBRES Y APELLIDO  ",
        "C.I:",
        "TELÉFONO",
        "VOCERIA",
        "FECHA DE NACIMIENTO",
        "EDAD",
        "SEXO",
        "PROFESION/OFICIO",
        "CORREO",
        "FECHAS DE ELECCION",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[2].height = 69.75

    data_align = Alignment(horizontal="center", vertical="center")
    data_font = Font(size=11, name="Calibri")
    for row_idx in range(3, 13):
        ws.row_dimensions[row_idx].height = 25.5
        cell_a = ws.cell(row=row_idx, column=1, value=row_idx - 2)
        cell_a.alignment = data_align
        cell_a.font = data_font
        for col_idx in range(2, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = data_align
            cell.font = data_font
            if col_idx == 11:
                cell.number_format = "@"
        for col_idx in (13, 18):
            ws.cell(row=row_idx, column=col_idx).number_format = "dd-mm-yyyy"

    for row_idx in range(3, 1001):
        ws.cell(row=row_idx, column=10).number_format = r"##\,###\,##0"
        celda_edad = ws.cell(row=row_idx, column=14)
        celda_edad.value = f'=IF(M{row_idx}="","",DATEDIF(M{row_idx},TODAY(),"Y"))'
        celda_edad.number_format = "0"

    col_widths = {
        "A": 8.28515625, "B": 14.5703125, "C": 14.0, "D": 16.85546875,
        "E": 15.7109375, "F": 26.85546875, "G": 45.42578125,
        "H": 25.7109375, "I": 26.7109375, "J": 17.0,
        "K": 17.5703125, "L": 18.140625, "M": 20.7109375,
        "N": 20.7109375, "O": 8.0, "P": 22.85546875,
        "Q": 20.42578125, "R": 23.7109375,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws_listas = wb.create_sheet("Catalogos")
    if ubicaciones_asic is None:
        ubicaciones_asic = [
            {"asic_nombre": asic, "municipio": municipio, "parroquia": parroquia}
            for asic, municipio, parroquia in UBICACIONES_ASIC_PREDETERMINADAS
        ]
    if not ubicaciones_asic:
        raise ValueError("No hay ubicaciones ASIC disponibles para incluir en la plantilla.")
    def agregar_lista(nombre, columna, opciones, rango, titulo):
        for fila, opcion in enumerate(opciones, start=2):
            ws_listas.cell(row=fila, column=columna, value=opcion)

        letra_columna = get_column_letter(columna)
        ws_listas.cell(row=1, column=columna, value=titulo)
        referencia = (
            f"'{ws_listas.title}'!${letra_columna}$2:"
            f"${letra_columna}${len(opciones) + 1}"
        )
        wb.defined_names.add(DefinedName(nombre, attr_text=referencia))

        validacion = DataValidation(type="list",         formula1=f"={nombre}", allow_blank=True)
        validacion.promptTitle = titulo
        validacion.prompt = "Seleccione una opción válida."
        validacion.add(rango)
        ws.add_data_validation(validacion)
        return validacion

    agregar_lista("ListaEstados", 1, VOCERIA_ESTADO_OPCIONES, "B3:B1000", "Estado")
    agregar_lista("ListaCentros", 4, VOCERIA_CENTRO_DISPENSADOR_OPCIONES, "G3:G1000", "Centro dispensador")
    agregar_lista("ListaVocerias", 5, VOCERIA_TIPO_OPCIONES, "L3:L1000", "Vocería")
    agregar_lista("ListaSexo", 6, VOCERIA_SEXO_OPCIONES, "O3:O1000", "Sexo")

    columna_asic = 7
    columna_municipios = 8
    nombres_asics = sorted({
        ubicacion["asic_nombre"] for ubicacion in ubicaciones_asic
    })
    municipios = sorted({
        ubicacion["municipio"] for ubicacion in ubicaciones_asic
    })
    parroquias_por_municipio = {}
    for ubicacion in ubicaciones_asic:
        municipio = ubicacion["municipio"]
        parroquias_por_municipio.setdefault(municipio, set()).add(
            ubicacion["parroquia"]
        )
    for fila, asic in enumerate(nombres_asics, start=1):
        ws_listas.cell(row=fila + 1, column=columna_asic, value=asic)
    for fila, municipio in enumerate(municipios, start=1):
        ws_listas.cell(row=fila + 1, column=columna_municipios, value=municipio)

    letra_columna_asic = get_column_letter(columna_asic)
    hoja_nombre = f"'{ws_listas.title}'!"
    ws_listas.cell(row=1, column=columna_asic, value="ASIC")
    ws_listas.cell(row=1, column=columna_municipios, value="Municipios")
    wb.defined_names.add(DefinedName(
        "CatalogoASIC",
        attr_text=f"{hoja_nombre}${letra_columna_asic}$2:${letra_columna_asic}${len(nombres_asics) + 1}",
    ))
    wb.defined_names.add(DefinedName(
        "CatalogoMunicipios",
        attr_text=(
            f"{hoja_nombre}${get_column_letter(columna_municipios)}$2:"
            f"${get_column_letter(columna_municipios)}${len(municipios) + 1}"
        ),
    ))
    columna_parroquias = columna_municipios + 1
    for indice, municipio in enumerate(municipios, start=1):
        columna = columna_parroquias + indice - 1
        opciones = sorted(parroquias_por_municipio[municipio])
        ws_listas.cell(row=1, column=columna, value=municipio)
        for fila, opcion in enumerate(opciones, start=2):
            ws_listas.cell(row=fila, column=columna, value=opcion)
        letra = get_column_letter(columna)
        wb.defined_names.add(DefinedName(
            f"Parroquia{indice}",
            attr_text=f"{hoja_nombre}${letra}$2:${letra}${len(opciones) + 1}",
        ))

    validacion_asic = DataValidation(
        type="list",
        formula1="=CatalogoASIC",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="ASIC no válido",
        error="Selecciona un ASIC de la lista.",
    )
    validacion_asic.add("E3:E1000")
    ws.add_data_validation(validacion_asic)
    validacion_asic.promptTitle = "ASIC"
    validacion_asic.prompt = "Seleccione el ASIC correspondiente."
    validacion_asic.showInputMessage = True

    dv_municipio = DataValidation(
        type="list",
        formula1="=CatalogoMunicipios",
        allow_blank=True,
    )
    dv_municipio.showErrorMessage = True
    dv_municipio.errorTitle = "Municipio no válido"
    dv_municipio.error = "Selecciona un municipio de la lista."
    dv_municipio.promptTitle = "Municipio"
    dv_municipio.prompt = "Seleccione un municipio."
    dv_municipio.showInputMessage = True
    dv_municipio.add("C3:C1000")
    ws.add_data_validation(dv_municipio)

    dv_parroquia = DataValidation(
        type="list",
        formula1='=INDIRECT("Parroquia"&MATCH(C3,CatalogoMunicipios,0))',
        allow_blank=True,
    )
    dv_parroquia.showErrorMessage = True
    dv_parroquia.errorTitle = "Parroquia no válida"
    dv_parroquia.error = "Selecciona una parroquia correspondiente al municipio."
    dv_parroquia.promptTitle = "Parroquia según municipio"
    dv_parroquia.prompt = "Seleccione primero el municipio."
    dv_parroquia.showInputMessage = True
    dv_parroquia.add("D3:D1000")
    ws.add_data_validation(dv_parroquia)

    dv_telefono = DataValidation(
        type="textLength",
        operator="between",
        formula1="11",
        formula2="12",
        allow_blank=True,
    )
    dv_telefono.promptTitle = ""
    dv_telefono.prompt = "xxxx-xxxxxxx"
    dv_telefono.showInputMessage = True
    dv_telefono.errorTitle = "Teléfono inválido"
    dv_telefono.error = "El teléfono debe tener entre 11 y 12 caracteres (ej: 0414-1234567)."
    dv_telefono.showErrorMessage = True
    dv_telefono.add("K3:K1000")
    ws.add_data_validation(dv_telefono)

    dv_fecha = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(1900,1,1)",
        formula2="DATE(9999,12,31)",
        allow_blank=True,
    )
    dv_fecha.promptTitle = ""
    dv_fecha.prompt = "DD-MM-AAAA"
    dv_fecha.showInputMessage = True
    dv_fecha.errorTitle = "Fecha no válida"
    dv_fecha.error = "Ingrese un valor de fecha válido (DD-MM-AAAA)."
    dv_fecha.showErrorMessage = True
    dv_fecha.add("M3:M1000")
    dv_fecha.add("R3:R1000")
    ws.add_data_validation(dv_fecha)

    dv_edad = DataValidation(
        type="whole",
        operator="between",
        formula1="18",
        formula2="99",
        allow_blank=True,
    )
    dv_edad.promptTitle = "Edad"
    dv_edad.prompt = "Ingrese una edad entera entre 18 y 99 años."
    dv_edad.errorTitle = "Edad no válida"
    dv_edad.error = "La edad debe ser un número entero entre 18 y 99."
    dv_edad.showErrorMessage = True
    dv_edad.add("N3:N1000")
    ws.add_data_validation(dv_edad)

    dv_ci = DataValidation(
        type="whole",
        operator="between",
        formula1="1000000",
        formula2="99999999",
        allow_blank=True,
    )
    dv_ci.promptTitle = "C.I."
    dv_ci.prompt = "Ingrese entre 7 y 8 dígitos. Se mostrarán comas cada tres dígitos."
    dv_ci.errorTitle = "C.I. no válida"
    dv_ci.error = "La C.I. debe ser un número entero de 7 u 8 dígitos."
    dv_ci.showErrorMessage = True
    dv_ci.add("J3:J1000")
    ws.add_data_validation(dv_ci)

    wb.save(output_path)
    return output_path


class ExcelService:
    """
    Servicio encargado de leer un archivo Excel y convertir sus filas
    en registros de actividad, con validación detallada por fila.
    """
    def __init__(self, id_usuario_registro: int):
        self.id_usuario_registro = id_usuario_registro
        self._actividad_dao = ActividadDAO()
        self._asic_dao = AsicDAO()
        self._municipio_dao = MunicipioDAO()
        self._parroquia_dao = ParroquiaDAO()
        self._tipo_dao = TipoActividadDAO()
        self._cache_asic: dict = {}
        self._cache_tipo: dict = {}
        self._cache_ubicaciones: dict = {}
        self._cache_municipios: set = set()

    def _validar_campos_fila(
        self, fila, idx, fecha, tema, comites_salud, participantes, id_asic
    ):
        errores = []
        valores_obligatorios = {
            "ESTADO": self._obtener_valor_columna(fila, idx, "estado"),
            "MUNICIPIO": self._obtener_valor_columna(fila, idx, "municipio"),
            "PARROQUIA": self._obtener_valor_columna(fila, idx, "parroquia"),
            "ESTABLECIMIENTO DE SALUD": self._obtener_valor_columna(fila, idx, "establecimiento_salud"),
            "NOMBRE DEL ESTABLECIMIENTO": self._obtener_valor_columna(fila, idx, "nombre_establecimiento_salud"),
        }
        for campo, valor in valores_obligatorios.items():
            if valor is None or not str(valor).strip():
                errores.append(f"El campo '{campo}' está vacío.")
        limites_texto = {
            "ESTADO": (valores_obligatorios["ESTADO"], 100),
            "MUNICIPIO": (valores_obligatorios["MUNICIPIO"], 100),
            "PARROQUIA": (valores_obligatorios["PARROQUIA"], 100),
            "ESTABLECIMIENTO DE SALUD": (valores_obligatorios["ESTABLECIMIENTO DE SALUD"], 150),
            "NOMBRE DEL ESTABLECIMIENTO": (valores_obligatorios["NOMBRE DEL ESTABLECIMIENTO"], 200),
        }
        for campo, (valor, maximo) in limites_texto.items():
            texto = str(valor or "").strip()
            if len(texto) > maximo or any(ord(caracter) < 32 for caracter in texto):
                errores.append(f"'{campo}' supera {maximo} caracteres o contiene controles.")
        estado = str(valores_obligatorios["ESTADO"] or "").strip()
        municipio = str(valores_obligatorios["MUNICIPIO"] or "").strip()
        establecimiento = str(valores_obligatorios["ESTABLECIMIENTO DE SALUD"] or "").strip()
        if estado and estado != "Táchira":
            errores.append("'ESTADO' no es válido.")
        if not getattr(self, "_cache_municipios", None):
            try:
                self._cache_municipios = {m["nombre"] for m in self._municipio_dao.obtener_todos()}
            except Exception:
                self._cache_municipios = set(MUNICIPIOS_TACHIRA)
        if municipio and municipio not in self._cache_municipios:
            errores.append("'MUNICIPIO' no es válido.")
        if id_asic and not any(
            ubicacion["municipio"] == municipio
            and ubicacion["parroquia"] == str(
                valores_obligatorios["PARROQUIA"] or ""
            ).strip()
            for ubicacion in self._ubicaciones_asic(id_asic)
        ):
            errores.append("'PARROQUIA' no corresponde al municipio seleccionado.")
        if establecimiento and establecimiento not in ACTIVIDAD_ESTABLECIMIENTO_OPCIONES:
            errores.append("'ESTABLECIMIENTO DE SALUD' no es válido.")

        if fecha is None:
            errores.append("'FECHA' no tiene formato de fecha válido (YYYY-MM-DD o DD/MM/YYYY).")
        elif fecha > date.today():
            errores.append("'FECHA' no puede ser futura.")

        if not tema:
            errores.append("El campo 'TEMA TRATADO' está vacío.")
        elif len(tema) > 150:
            errores.append("'TEMA TRATADO' no puede superar los 150 caracteres.")

        if comites_salud is None or not 0 <= comites_salud <= 99:
            errores.append(
                "'CANTIDAD DE COMITÉS DE SALUD' debe ser un número entero entre 0 y 99; "
                "ingrese la cantidad, no el nombre del comité."
            )

        participantes_por_categoria = []
        for campo in PARTICIPANTES_CATEGORIAS:
            valor = self._obtener_valor_columna(fila, idx, campo)
            if valor is None or not str(valor).strip():
                participantes_por_categoria.append(0)
                continue
            cantidad = self._parse_int(valor)
            if cantidad is None or not 0 <= cantidad <= 99999:
                errores.append(
                    f"'{campo.replace('_', ' ').upper()}' debe ser un número entero entre 0 y 99999."
                )
                participantes_por_categoria.append(0)
                continue
            participantes_por_categoria.append(cantidad)

        hay_desglose = any(
            valor is not None and str(valor).strip() != ""
            for campo in PARTICIPANTES_CATEGORIAS
            for valor in [self._obtener_valor_columna(fila, idx, campo)]
        )
        if hay_desglose and participantes is not None and sum(participantes_por_categoria) != participantes:
            errores.append("'TOTAL DE ASISTENCIA' debe coincidir con la suma de las categorías de participantes.")

        return errores

    def _ubicaciones_asic(self, id_asic):
        if id_asic not in self._cache_ubicaciones:
            self._cache_ubicaciones[id_asic] = self._parroquia_dao.listar_por_asic(
                id_asic
            )
        return self._cache_ubicaciones[id_asic]

    def previsualizar(self, ruta_archivo: str) -> dict:
        validar_archivo_excel(ruta_archivo)
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
        ws = wb.active
        raw_header_row = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
        cabecera_raw = [str(c).strip() if c is not None else "" for c in raw_header_row]
        cabecera_normalizada = [self._normalizar_encabezado(c) for c in cabecera_raw]
        idx = {cab: i for i, cab in enumerate(cabecera_normalizada) if cab}
        data_start = 4
        raw_row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        for i, val in enumerate(raw_row2):
            if i < len(cabecera_raw) and not cabecera_raw[i] and val is not None:
                normalizado = self._normalizar_encabezado(str(val))
                if normalizado:
                    idx[normalizado] = i
                    cabecera_raw[i] = str(val).strip()
        if not any(col in idx for col in COLUMNAS_REQUERIDAS):
            raw_header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            cabecera_raw = [str(c).strip() if c is not None else "" for c in raw_header_row]
            cabecera_normalizada = [self._normalizar_encabezado(c) for c in cabecera_raw]
            idx = {cab: i for i, cab in enumerate(cabecera_normalizada) if cab}
            data_start = 2
        faltantes = []
        for col in COLUMNAS_REQUERIDAS:
            if col in idx:
                continue
            if col == "participantes" and ("total_asistencia" in idx or any(cf in idx for cf in PARTICIPANTES_CATEGORIAS)):
                continue
            faltantes.append(col)
        if faltantes:
            raise ValueError(
                f"El archivo no tiene las columnas requeridas: {', '.join(faltantes)}.\n"
                f"Columnas detectadas: {', '.join(cabecera_raw)}"
            )
        filas_validas = []
        filas_error   = []
        for numero_fila, fila in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            if all(v is None or str(v).strip() == "" for v in fila):
                continue
            datos_crudos = {col: fila[idx[col]] if col in idx and idx[col] < len(fila) else None for col in COLUMNAS_REQUERIDAS}
            errores = []
            asic_nombre  = str(datos_crudos["asic"]).strip()          if datos_crudos["asic"]          else ""
            tipo_nombre  = str(datos_crudos["tipo_actividad"]).strip() if datos_crudos["tipo_actividad"] else ""
            tema = str(datos_crudos.get("tema") or "").strip()
            id_asic = self._resolver_asic(asic_nombre) if asic_nombre else None
            comites_salud = self._parse_int(self._obtener_valor_columna(fila, idx, "comites_de_salud"))
            try:
                participantes = self._calcular_total_asistencia(fila, idx, datos_crudos)
                if participantes is None or not 0 <= participantes <= 899991:
                    raise ValueError()
            except (TypeError, ValueError):
                participantes = None
                errores.append("'TOTAL DE ASISTENCIA' debe ser un entero entre 0 y 899991.")
            fecha_actividad = self._parsear_fecha(datos_crudos.get("fecha"))
            errores.extend(self._validar_campos_fila(
                fila, idx, fecha_actividad, tema, comites_salud, participantes,
                id_asic,
            ))
            if not asic_nombre:
                errores.append("El campo 'ASIC' está vacío.")
            elif id_asic is None:
                errores.append(f"ASIC no encontrado en BD: «{asic_nombre}».")
            if not tipo_nombre:
                errores.append("El campo 'ACTIVIDAD DESARROLLADA' está vacío.")
            id_tipo = None
            if tipo_nombre:
                id_tipo = self._resolver_tipo(tipo_nombre)
                if id_tipo is None:
                    errores.append(f"Tipo de actividad no encontrado en BD: «{tipo_nombre}»." )
            if id_asic and id_tipo and fecha_actividad and tema and not errores:
                es_duplicado = self._actividad_dao.verificar_duplicado(
                    id_asic, id_tipo, fecha_actividad, tema=tema
                )
                if es_duplicado:
                    errores.append("Duplicado: ya existe una actividad con el mismo ASIC, tipo, fecha y tema.")
            fila_procesada = {
                "numero_fila": numero_fila,
                "id_asic": id_asic,
                "id_tipo": id_tipo,
                "fecha": fecha_actividad,
                "estado": self._obtener_valor_columna(fila, idx, "estado"),
                "municipio": self._obtener_valor_columna(fila, idx, "municipio"),
                "parroquia": self._obtener_valor_columna(fila, idx, "parroquia"),
                "establecimiento_salud": self._obtener_valor_columna(fila, idx, "establecimiento_salud"),
                "nombre_establecimiento": self._obtener_valor_columna(fila, idx, "nombre_establecimiento_salud"),
                "comites_salud": comites_salud,
                "tema": tema,
                "medicos": self._parse_int(self._obtener_valor_columna(fila, idx, "medicos")),
                "promotores_salud": self._parse_int(self._obtener_valor_columna(fila, idx, "promotores_en_salud")),
                "voceros_comite_salud": self._parse_int(self._obtener_valor_columna(fila, idx, "voceros_de_comite_de_salud")),
                "vocero_consejo_direccion": self._parse_int(self._obtener_valor_columna(fila, idx, "vocero_del_consejo_de_direccion_del_asic")),
                "coordinador_docencia": self._parse_int(self._obtener_valor_columna(fila, idx, "coordinador_de_docencia_e_investigacion_del_asic")),
                "autoridades_presentes": self._parse_int(self._obtener_valor_columna(fila, idx, "autoridades_presentes")),
                "otros_actores": self._parse_int(self._obtener_valor_columna(fila, idx, "otros_actores_involucrados")),
                "personas_comunidad": self._parse_int(self._obtener_valor_columna(fila, idx, "cantidad_personas_comunidad")),
                "estudiantes_aspirantes": self._parse_int(self._obtener_valor_columna(fila, idx, "estudiantes_aspirantes")),
                "participantes": participantes,
                "datos_crudos": datos_crudos,
            }
            if errores:
                filas_error.append({**fila_procesada, "errores": errores})
            else:
                filas_validas.append(fila_procesada)
        return {
            "columnas": cabecera_raw,
            "filas_validas": filas_validas,
            "filas_error": filas_error,
            "total_filas": len(filas_validas) + len(filas_error),
        }
    def importar(self, filas_validas: list) -> dict:
        import db

        auth = db.AuthController()
        usuario = auth.obtener_usuario_actual()
        if not auth.usuario_puede_escribir(usuario):
            raise PermissionError(auth.mensaje_sin_permiso())
        if not filas_validas:
            return {"insertadas": 0, "errores_al_insertar": []}

        conn = db.get_connection()
        fila_actual = None
        try:
            for fila in filas_validas:
                fila_actual = fila
                self._actividad_dao.crear(
                    id_asic=fila["id_asic"],
                    id_tipo=fila["id_tipo"],
                    id_usuario_registro=self.id_usuario_registro,
                    fecha=fila.get("fecha"),
                    estado=fila.get("estado") or "",
                    municipio=fila.get("municipio") or "",
                    parroquia=fila.get("parroquia") or "",
                    establecimiento_salud=fila.get("establecimiento_salud") or "",
                    nombre_establecimiento=fila.get("nombre_establecimiento") or "",
                    comites_salud=fila.get("comites_salud"),
                    tema=fila.get("tema") or "",
                    medicos=fila.get("medicos") or 0,
                    promotores_salud=fila.get("promotores_salud") or 0,
                    voceros_comite_salud=fila.get("voceros_comite_salud") or 0,
                    vocero_consejo_direccion=fila.get("vocero_consejo_direccion") or 0,
                    coordinador_docencia=fila.get("coordinador_docencia") or 0,
                    autoridades_presentes=fila.get("autoridades_presentes") or 0,
                    otros_actores=fila.get("otros_actores") or 0,
                    personas_comunidad=fila.get("personas_comunidad") or 0,
                    estudiantes_aspirantes=fila.get("estudiantes_aspirantes") or 0,
                    participantes=fila["participantes"],
                    connection=conn,
                )
            conn.commit()
            return {"insertadas": len(filas_validas), "errores_al_insertar": []}
        except Exception as exc:
            conn.rollback()
            return {
                "insertadas": 0,
                "errores_al_insertar": [{
                    "numero_fila": fila_actual.get("numero_fila") if fila_actual else "-",
                    "error": str(exc),
                }],
            }
        finally:
            conn.close()

    def _normalizar_encabezado(self, texto):
        if texto is None:
            return ""
        valor = str(texto).strip().lower()
        valor = unicodedata.normalize("NFKD", valor)
        valor = "".join(ch for ch in valor if not unicodedata.combining(ch))
        valor = re.sub(r"[^a-z0-9]+", " ", valor)
        valor = " ".join(valor.split())
        return HEADER_ALIASES.get(valor, valor)

    def _obtener_valor_columna(self, fila, idx, campo):
        if campo not in idx or idx[campo] >= len(fila):
            return None
        return fila[idx[campo]]

    def _parse_int(self, valor):
        if valor is None:
            return None
        if isinstance(valor, bool):
            return None
        if isinstance(valor, (int, float)) and float(valor).is_integer():
            return int(valor)
        try:
            texto = str(valor).strip()
            if texto == "":
                return None
            return int(texto)
        except (TypeError, ValueError):
            return None

    def _calcular_total_asistencia(self, fila, idx, datos_crudos):
        total = self._obtener_valor_columna(fila, idx, "total_asistencia")
        total_num = self._parse_int(total)
        if total_num is not None:
            return total_num
        suma = 0
        encontrado = False
        for campo in PARTICIPANTES_CATEGORIAS:
            valor = self._obtener_valor_columna(fila, idx, campo)
            if valor is None or str(valor).strip() == "":
                continue
            encontrado = True
            valor_num = self._parse_int(valor)
            if valor_num is None:
                return None
            suma += valor_num
        if encontrado:
            return suma
        return self._parse_int(datos_crudos.get("participantes"))
    def _parsear_fecha(self, valor):
        """Acepta objetos date/datetime o strings YYYY-MM-DD / DD/MM/YYYY."""
        if valor is None:
            return None
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        texto = str(valor).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(texto, fmt).date()
            except ValueError:
                pass
        return None
    def _resolver_asic(self, nombre: str):
        """Busca el id_asic por nombre (case-insensitive). Usa caché."""
        key = nombre.lower().strip()
        if key in self._cache_asic:
            return self._cache_asic[key]
        for asic in self._asic_dao.obtener_todos():
            if asic["nombre"].lower().strip() == key:
                self._cache_asic[key] = asic["id_asic"]
                return asic["id_asic"]
        self._cache_asic[key] = None
        return None
    def _resolver_tipo(self, nombre: str):
        """Busca el id_tipo por nombre (case-insensitive). Usa caché."""
        key = nombre.lower().strip()
        if key in self._cache_tipo:
            return self._cache_tipo[key]
        for tipo in self._tipo_dao.obtener_todos():
            if tipo["nombre"].lower().strip() == key:
                self._cache_tipo[key] = tipo["id_tipo"]
                return tipo["id_tipo"]
        self._cache_tipo[key] = None
        return None


class ExcelServiceVoceria:
    VOCERIA_HEADERS = [
        "nro",
        "estado",
        "municipio",
        "parroquia",
        "asic",
        "nombre_comite",
        "centro_dispensador",
        "nombre_centro",
        "nombres",
        "ci",
        "telefono",
        "voceria",
        "fecha_nacimiento",
        "edad",
        "sexo",
        "profesion",
        "correo",
        "fechas_eleccion",
    ]

    HEADER_ALIASES_VOCERIA = {
        "nro": "nro",
        "n": "nro",
        "numero": "nro",
        "estado": "estado",
        "municipio": "municipio",
        "parroquia": "parroquia",
        "asic": "asic",
        "nombre del comite": "nombre_comite",
        "nombre del comité": "nombre_comite",
        "indique que centro dispensador de salud es sri hospital cat cdi clinica popular consultorio medico popular ambulatorio u otro": "centro_dispensador",
        "indique que centro dispensador de salud es": "centro_dispensador",
        "centro dispensador": "centro_dispensador",
        "nombre del centro dispensador de salud": "nombre_centro",
        "nombre del centro": "nombre_centro",
        "nombres y apellido": "nombres",
        "nombres": "nombres",
        "c i": "ci",
        "c.i.": "ci",
        "ci": "ci",
        "telefono": "telefono",
        "teléfono": "telefono",
        "voceria": "voceria",
        "vocería": "voceria",
        "fecha de nacimiento": "fecha_nacimiento",
        "edad": "edad",
        "sexo": "sexo",
        "profesion oficio": "profesion",
        "profesion/oficio": "profesion",
        "profesion": "profesion",
        "profesión": "profesion",
        "correo": "correo",
        "fechas de eleccion": "fechas_eleccion",
        "fechas de elección": "fechas_eleccion",
        "fecha de eleccion": "fechas_eleccion",
        "fecha de elección": "fechas_eleccion",
    }

    CAMPOS_OBLIGATORIOS = [
        "estado", "municipio", "parroquia", "asic", "nombre_comite",
        "centro_dispensador", "nombre_centro", "nombres", "ci",
        "telefono", "voceria", "fecha_nacimiento", "edad",
        "sexo", "profesion", "correo", "fechas_eleccion"
    ]

    def __init__(self, id_usuario_registro: int):
        self.id_usuario_registro = id_usuario_registro
        self._asic_dao = AsicDAO()
        self._municipio_dao = MunicipioDAO()
        self._parroquia_dao = ParroquiaDAO()
        self._tipo_voceria_dao = TipoVoceriaDAO()
        self._cache_asic: dict = {}
        self._cache_tipo_voceria: dict = {}
        self._cache_ubicaciones: dict = {}
        self._cache_municipios: set = set()

    def previsualizar(self, ruta_archivo: str) -> dict:
        import openpyxl
        from datetime import date, datetime, timedelta

        validar_archivo_excel(ruta_archivo)
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
        ws = wb.active

        raw_header_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        cabecera_raw = [str(c).strip() if c is not None else "" for c in raw_header_row]
        cabecera_normalizada = [self._normalizar_encabezado(c) for c in cabecera_raw]
        idx = {cab: i for i, cab in enumerate(cabecera_normalizada) if cab}

        claves_esperadas = {"estado", "municipio", "parroquia", "asic", "nombre_comite",
                           "centro_dispensador", "nombre_centro", "nombres", "ci",
                           "telefono", "voceria", "fecha_nacimiento", "edad",
                           "sexo", "profesion", "correo", "fechas_eleccion"}
        if not claves_esperadas.intersection(idx.keys()):
            raw_header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            cabecera_raw = [str(c).strip() if c is not None else "" for c in raw_header_row]
            cabecera_normalizada = [self._normalizar_encabezado(c) for c in cabecera_raw]
            idx = {cab: i for i, cab in enumerate(cabecera_normalizada) if cab}
            data_start = 2
        else:
            data_start = 3

        if not idx:
            raise ValueError("No se detectaron encabezados válidos en el archivo.")

        filas_por_comite: dict = {}
        errores_globales = []

        for numero_fila, fila in enumerate(ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
            if all(v is None or str(v).strip() == "" for v in fila):
                continue

            datos_crudos = self._extraer_datos_fila(fila, idx, cabecera_raw)
            errores_fila = self._validar_fila(datos_crudos, numero_fila)

            nombre_comite = (datos_crudos.get("nombre_comite") or "").strip()
            if not nombre_comite:
                errores_fila.append("El campo 'NOMBRE DEL COMITÉ' está vacío.")

            if errores_fila:
                errores_globales.append({
                    "numero_fila": numero_fila,
                    "nombre_comite": nombre_comite or "(sin comité)",
                    "errores": errores_fila,
                    "datos_crudos": datos_crudos
                })
            else:
                if nombre_comite not in filas_por_comite:
                    filas_por_comite[nombre_comite] = []
                filas_por_comite[nombre_comite].append({
                    "numero_fila": numero_fila,
                    "datos": datos_crudos
                })

        comites_validos = []
        comites_error = []

        for nombre_comite, filas in filas_por_comite.items():
            errores_comite = self._validar_comite(nombre_comite, filas)
            if errores_comite:
                for f in filas:
                    comites_error.append({
                        "numero_fila": f["numero_fila"],
                        "nombre_comite": nombre_comite,
                        "errores": errores_comite,
                        "datos_crudos": f["datos"]
                    })
            else:
                comites_validos.append({
                    "nombre_comite": nombre_comite,
                    "filas": filas
                })

        return {
            "columnas": cabecera_raw,
            "comites_validos": comites_validos,
            "comites_error": comites_error,
            "filas_error_sin_comite": errores_globales,
            "total_filas_procesadas": sum(len(c["filas"]) for c in comites_validos) + len(comites_error) + len(errores_globales),
            "total_comites_validos": len(comites_validos),
            "total_comites_error": len(set(c["nombre_comite"] for c in comites_error)),
        }

    def _normalizar_encabezado(self, texto):
        if texto is None:
            return ""
        valor = str(texto).strip().lower()
        import unicodedata
        valor = unicodedata.normalize("NFKD", valor)
        valor = "".join(ch for ch in valor if not unicodedata.combining(ch))
        import re
        valor = re.sub(r"[^a-z0-9]+", " ", valor)
        valor = " ".join(valor.split())
        return self.HEADER_ALIASES_VOCERIA.get(valor, valor)

    def _extraer_datos_fila(self, fila, idx, cabecera_raw):
        datos = {}
        for campo in self.VOCERIA_HEADERS:
            if campo in idx and idx[campo] < len(fila):
                valor = fila[idx[campo]]
                if campo == "telefono":
                    datos[campo] = self._normalizar_telefono_excel(valor)
                elif campo in ("fecha_nacimiento", "fechas_eleccion") and isinstance(valor, (date, datetime)):
                    datos[campo] = valor.strftime("%Y-%m-%d")
                elif campo == "ci" and isinstance(valor, (int, float)) and not isinstance(valor, bool):
                    datos[campo] = str(int(valor)) if float(valor).is_integer() else str(valor).strip()
                elif campo == "edad" and isinstance(valor, (int, float)) and not isinstance(valor, bool):
                    datos[campo] = str(int(valor)) if float(valor).is_integer() else str(valor).strip()
                else:
                    datos[campo] = str(valor).strip() if valor is not None else ""
            else:
                datos[campo] = ""

        if not datos.get("edad"):
            fecha_nacimiento = self._parsear_fecha(datos.get("fecha_nacimiento"))
            if fecha_nacimiento:
                hoy = date.today()
                edad = hoy.year - fecha_nacimiento.year - (
                    (hoy.month, hoy.day) < (fecha_nacimiento.month, fecha_nacimiento.day)
                )
                datos["edad"] = str(edad)
        return datos

    def _normalizar_telefono_excel(self, valor):
        if valor is None:
            return ""

        es_numerico = isinstance(valor, (int, float)) and not isinstance(valor, bool)
        if isinstance(valor, float):
            if not valor.is_integer():
                return str(valor).strip()
            texto = str(int(valor))
        else:
            texto = str(valor).strip()

        if es_numerico and len(texto) == 10 and texto.startswith(("412", "414", "416", "424", "426")):
            texto = "0" + texto

        texto_digitos = texto[1:] if texto.startswith("+") else texto
        separadores = " ()-./"
        if all(caracter.isdigit() or caracter.isspace() or caracter in separadores for caracter in texto_digitos):
            return "".join(caracter for caracter in texto_digitos if caracter.isdigit())
        return texto

    def _validar_fila(self, datos, numero_fila):
        errores = []
        for campo in self.CAMPOS_OBLIGATORIOS:
            valor = (datos.get(campo) or "").strip()
            if not valor or valor.lower() in ("aaaa-mm-dd", "edad", "seleccionar...", "(sin opciones disponibles)"):
                nombre_campo = campo.replace("_", " ").upper()
                errores.append(f"Fila {numero_fila}: El campo '{nombre_campo}' es obligatorio.")
        correo = (datos.get("correo") or "").strip()
        if correo and not validar_correo(correo):
            errores.append(f"Fila {numero_fila}: Correo electrónico no válido.")

        nombres = (datos.get("nombres") or "").strip()
        if nombres and not 5 <= len(nombres) <= 40:
            errores.append(f"Fila {numero_fila}: Nombres y Apellidos debe tener entre 5 y 40 caracteres.")
        if nombres and any(caracter.isdigit() or ord(caracter) < 32 for caracter in nombres):
            errores.append(f"Fila {numero_fila}: Nombres y Apellidos no puede contener números ni controles.")

        ci = (datos.get("ci") or "").strip()
        ci_digitos = ci.replace(",", "")
        formato_ci_valido = bool(re.fullmatch(r"(?:\d{7,8}|\d{1,2}(?:,\d{3}){2})", ci))
        if ci and (not formato_ci_valido or not 7 <= len(ci_digitos) <= 8):
            errores.append(f"Fila {numero_fila}: La C.I. debe tener entre 7 y 8 dígitos.")

        profesion = (datos.get("profesion") or "").strip()
        if profesion and not 3 <= len(profesion) <= 40:
            errores.append(f"Fila {numero_fila}: Profesión/Oficio debe tener entre 3 y 40 caracteres.")

        telefono = (datos.get("telefono") or "").strip()
        if telefono and (not telefono.isdigit() or not 11 <= len(telefono) <= 12):
            errores.append(f"Fila {numero_fila}: El teléfono debe tener entre 11 y 12 dígitos.")

        edad = (datos.get("edad") or "").strip()
        if edad and (not edad.isdigit() or len(edad) > 2 or not 18 <= int(edad) <= 99):
            errores.append(f"Fila {numero_fila}: La edad debe estar entre 18 y 99 años.")

        sexo = (datos.get("sexo") or "").strip()
        if sexo and sexo not in ("M", "F"):
            errores.append(f"Fila {numero_fila}: Sexo debe ser M o F.")

        limites_generales = {
            "estado": 100, "municipio": 100, "parroquia": 100,
            "centro_dispensador": 100, "nombre_centro": 200, "nombre_comite": 200,
        }
        for campo, maximo in limites_generales.items():
            valor = (datos.get(campo) or "").strip()
            if valor and (len(valor) > maximo or any(ord(caracter) < 32 for caracter in valor)):
                errores.append(f"Fila {numero_fila}: '{campo.replace('_', ' ').upper()}' supera {maximo} caracteres o contiene controles.")

        estado = (datos.get("estado") or "").strip()
        if estado and estado not in VOCERIA_ESTADO_OPCIONES:
            errores.append(f"Fila {numero_fila}: Estado no válido.")
        municipio = (datos.get("municipio") or "").strip()
        if not getattr(self, "_cache_municipios", None):
            try:
                self._cache_municipios = {m["nombre"] for m in self._municipio_dao.obtener_todos()}
            except Exception:
                self._cache_municipios = set(MUNICIPIOS_TACHIRA)
        if municipio and municipio not in self._cache_municipios:
            errores.append(f"Fila {numero_fila}: Municipio no válido.")
        asic_nombre = (datos.get("asic") or "").strip()
        id_asic = self._resolver_asic(asic_nombre) if asic_nombre else None
        if asic_nombre and id_asic is None:
            errores.append(f"Fila {numero_fila}: ASIC no encontrado: «{asic_nombre}».")
        elif id_asic and not any(
            ubicacion["municipio"] == municipio
            and ubicacion["parroquia"] == (datos.get("parroquia") or "").strip()
            for ubicacion in self._ubicaciones_asic(id_asic)
        ):
            errores.append(
                f"Fila {numero_fila}: La parroquia no corresponde al municipio seleccionado."
            )
        centro_dispensador = (datos.get("centro_dispensador") or "").strip()
        if centro_dispensador and centro_dispensador not in VOCERIA_CENTRO_DISPENSADOR_OPCIONES:
            errores.append(f"Fila {numero_fila}: Centro dispensador no válido.")

        hoy = date.today()
        fecha_nacimiento = self._parsear_fecha(datos.get("fecha_nacimiento"))
        if datos.get("fecha_nacimiento") and not fecha_nacimiento:
            errores.append(f"Fila {numero_fila}: Fecha de Nacimiento no válida.")
        elif fecha_nacimiento:
            edad_por_fecha = hoy.year - fecha_nacimiento.year - (
                (hoy.month, hoy.day) < (fecha_nacimiento.month, fecha_nacimiento.day)
            )
            if not 18 <= edad_por_fecha <= 99:
                errores.append(f"Fila {numero_fila}: La persona debe tener entre 18 y 99 años según su Fecha de Nacimiento.")
            if edad.isdigit() and int(edad) != edad_por_fecha:
                errores.append(f"Fila {numero_fila}: La edad no coincide con la Fecha de Nacimiento.")

        fecha_eleccion = self._parsear_fecha(datos.get("fechas_eleccion"))
        if datos.get("fechas_eleccion") and not fecha_eleccion:
            errores.append(f"Fila {numero_fila}: Fecha de Elección no válida.")
        elif fecha_eleccion and fecha_eleccion > hoy:
            errores.append(f"Fila {numero_fila}: La Fecha de Elección no puede ser futura.")
        elif fecha_eleccion and fecha_eleccion < hoy - timedelta(days=730):
            errores.append(f"Fila {numero_fila}: La Fecha de Elección no puede tener más de 2 años.")
        return errores

    def _ubicaciones_asic(self, id_asic):
        if id_asic not in self._cache_ubicaciones:
            self._cache_ubicaciones[id_asic] = self._parroquia_dao.listar_por_asic(
                id_asic
            )
        return self._cache_ubicaciones[id_asic]

    def _validar_comite(self, nombre_comite, filas):
        errores = []

        if len(filas) < 7:
            errores.append(f"El comité '{nombre_comite}' debe tener mínimo 7 vocerías (actual: {len(filas)}).")
        if len(filas) > 10:
            errores.append(f"El comité '{nombre_comite}' debe tener máximo 10 vocerías (actual: {len(filas)}).")

        vocerias_vistas = set()
        for f in filas:
            datos = f["datos"]
            voceria = (datos.get("voceria") or "").strip().lower()
            if voceria in vocerias_vistas:
                errores.append(f"Comité '{nombre_comite}': Vocería duplicada '{datos.get('voceria')}'.")
            vocerias_vistas.add(voceria)

            fecha_eleccion_str = (datos.get("fechas_eleccion") or "").strip()
            fecha_eleccion = self._parsear_fecha(fecha_eleccion_str)
            if fecha_eleccion:
                hace_dos_anos = date.today() - timedelta(days=730)
                if fecha_eleccion < hace_dos_anos:
                    errores.append(
                        f"Comité '{nombre_comite}', fila {f['numero_fila']}: "
                        f"La fecha de elección ({fecha_eleccion_str}) tiene más de 2 años de antigüedad."
                    )

        return errores

    def _parsear_fecha(self, valor):
        if valor is None:
            return None
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        texto = str(valor).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(texto, fmt).date()
            except ValueError:
                pass
        return None

    def importar(self, comites_validos: list) -> dict:
        import db

        auth = db.AuthController()
        usuario = auth.obtener_usuario_actual()
        if not auth.usuario_puede_escribir(usuario):
            raise PermissionError(auth.mensaje_sin_permiso())
        voceria_ctrl = db.VoceriaController()

        if not comites_validos:
            return {
                "comites_insertados": 0,
                "voceros_insertados": 0,
                "errores_al_insertar": [],
            }

        conn = db.get_connection()
        comite_actual = None
        comites_insertados = 0
        voceros_insertados = 0
        try:
            for comite in comites_validos:
                comite_actual = comite
                primera_fila = comite["filas"][0]["datos"]
                id_asic = self._resolver_asic(primera_fila.get("asic", ""))
                if not id_asic:
                    raise ValueError(f"ASIC no encontrado: {primera_fila.get('asic')}")

                id_tipo_voceria = self._resolver_tipo_voceria(primera_fila.get("voceria", ""))
                if not id_tipo_voceria:
                    tipos = self._tipo_voceria_dao.obtener_todos()
                    if tipos:
                        id_tipo_voceria = tipos[0]["id_tipo"]
                    else:
                        raise ValueError("No hay tipos de vocería configurados en el sistema.")

                items = []
                for f in comite["filas"]:
                    d = f["datos"]
                    item = {
                        "nombres": d.get("nombres", ""),
                        "ci": d.get("ci", ""),
                        "telefono": d.get("telefono", ""),
                        "fecha_nacimiento": d.get("fecha_nacimiento", ""),
                        "edad": d.get("edad", ""),
                        "id_tipo_voceria": self._resolver_tipo_voceria(d.get("voceria", "")) or id_tipo_voceria,
                        "sexo": d.get("sexo", ""),
                        "profesion": d.get("profesion", ""),
                        "correo": d.get("correo", ""),
                        "fechas_eleccion": d.get("fechas_eleccion", ""),
                    }
                    items.append(item)

                ok, msg = voceria_ctrl.guardar_voceria(
                    id_asic=id_asic,
                    id_tipo=id_tipo_voceria,
                    estado=primera_fila.get("estado", ""),
                    municipio=primera_fila.get("municipio", ""),
                    parroquia=primera_fila.get("parroquia", ""),
                    centro_dispensador=primera_fila.get("centro_dispensador", ""),
                    nombre_centro=primera_fila.get("nombre_centro", ""),
                    nombre_comite=primera_fila.get("nombre_comite", ""),
                    items=items,
                    connection=conn,
                )
                if not ok:
                    raise Exception(msg)
                comites_insertados += 1
                voceros_insertados += len(items)

            conn.commit()
            return {
                "comites_insertados": comites_insertados,
                "voceros_insertados": voceros_insertados,
                "errores_al_insertar": [],
            }
        except Exception as exc:
            conn.rollback()
            return {
                "comites_insertados": 0,
                "voceros_insertados": 0,
                "errores_al_insertar": [{
                    "nombre_comite": comite_actual.get("nombre_comite", "(desconocido)") if comite_actual else "(desconocido)",
                    "error": str(exc)
                }],
            }
        finally:
            conn.close()

    def _resolver_asic(self, nombre: str):
        key = nombre.lower().strip()
        if key in self._cache_asic:
            return self._cache_asic[key]
        for asic in self._asic_dao.obtener_todos():
            if asic["nombre"].lower().strip() == key:
                self._cache_asic[key] = asic["id_asic"]
                return asic["id_asic"]
        self._cache_asic[key] = None
        return None

    def _resolver_tipo_voceria(self, nombre: str):
        key = nombre.lower().strip()
        if key in self._cache_tipo_voceria:
            return self._cache_tipo_voceria[key]
        for tipo in self._tipo_voceria_dao.obtener_todos():
            if tipo["nombre"].lower().strip() == key:
                self._cache_tipo_voceria[key] = tipo["id_tipo"]
                return tipo["id_tipo"]
        self._cache_tipo_voceria[key] = None
        return None
