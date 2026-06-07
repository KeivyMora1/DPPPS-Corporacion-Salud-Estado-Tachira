# ============================================================
# services/pdf_service.py
# Generación de informes consolidados en PDF con ReportLab.
#
# Soporta dos métodos de exportación:
#   - generar_pdf(datos, ruta_salida)  → genera el PDF en disco
#   - generar_excel(datos, ruta_salida) → genera el Excel de detalle
# ============================================================

import os
from datetime import datetime

# ReportLab
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    HRFlowable, PageBreak, Paragraph, SimpleDocTemplate,
    Spacer, Table, TableStyle,
)

# openpyxl para exportación Excel
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import settings

# ── Paleta de colores ──────────────────────────────────────────────────
AZUL_OSCURO  = colors.HexColor("#0f3460")
AZUL_MEDIO   = colors.HexColor("#16213e")
ROJO_ACENTO  = colors.HexColor("#e94560")
VERDE        = colors.HexColor("#05c46b")
GRIS_CLARO   = colors.HexColor("#f4f6f8")
GRIS_BORDE   = colors.HexColor("#cccccc")
BLANCO       = colors.white
NEGRO        = colors.black


class PDFService:
    """Genera el PDF del informe consolidado usando ReportLab."""

    # ------------------------------------------------------------------ 
    # PDF
    # ------------------------------------------------------------------

    def generar_pdf(self, datos: dict, ruta_salida: str) -> str:
        """
        Genera el PDF del informe y lo guarda en ruta_salida.
        :param datos: dict devuelto por ConsolidacionService.obtener_informe_completo()
                      o generar_previo() enriquecido con 'titulo' y 'observaciones'.
        :param ruta_salida: ruta completa del archivo .pdf de destino.
        :returns: ruta_salida (para comodidad del llamador).
        """
        os.makedirs(os.path.dirname(ruta_salida) if os.path.dirname(ruta_salida) else ".", exist_ok=True)

        doc = SimpleDocTemplate(
            ruta_salida,
            pagesize=landscape(A4),
            leftMargin=2*cm, rightMargin=2*cm,
            topMargin=2*cm,  bottomMargin=2*cm,
            title=datos.get("titulo", "Informe Consolidado"),
        )

        estilos = self._estilos()
        historia = []

        # ── Encabezado institucional ─────────────────────────────────
        historia += self._encabezado(datos, estilos)
        historia.append(Spacer(1, 0.4*cm))
        historia.append(HRFlowable(width="100%", thickness=2, color=AZUL_OSCURO))
        historia.append(Spacer(1, 0.3*cm))

        # ── Tarjetas de resumen ──────────────────────────────────────
        historia += self._tarjetas_resumen(datos, estilos)
        historia.append(Spacer(1, 0.5*cm))

        # ── Tabla resumen por ASIC ───────────────────────────────────
        resumen_asic = datos.get("resumen_por_asic") or datos.get("resumen_por_asic", [])
        if resumen_asic:
            historia += self._tabla_resumen_asic(resumen_asic, estilos)
            historia.append(Spacer(1, 0.5*cm))

        # ── Tabla resumen por tipo ───────────────────────────────────
        resumen_tipo = datos.get("resumen_por_tipo", [])
        if resumen_tipo:
            historia += self._tabla_resumen_tipo(resumen_tipo, estilos)
            historia.append(Spacer(1, 0.5*cm))

        # ── Detalle de actividades ───────────────────────────────────
        actividades = datos.get("actividades", [])
        if actividades:
            historia.append(PageBreak())
            historia += self._tabla_detalle(actividades, estilos)

        # ── Observaciones ────────────────────────────────────────────
        obs = datos.get("observaciones", "")
        if obs and obs.strip():
            historia.append(Spacer(1, 0.5*cm))
            historia.append(Paragraph("Observaciones", estilos["subtitulo"]))
            historia.append(Paragraph(obs, estilos["normal"]))

        # ── Pie de página con fecha de generación ────────────────────
        historia.append(Spacer(1, 0.5*cm))
        historia.append(HRFlowable(width="100%", thickness=1, color=GRIS_BORDE))
        historia.append(Paragraph(
            f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')} — "
            f"{settings.APP_TITLE} v{settings.APP_VERSION}",
            estilos["pie"],
        ))

        doc.build(historia)
        return ruta_salida

    # ------------------------------------------------------------------ 
    # Excel de detalle
    # ------------------------------------------------------------------

    def generar_excel(self, datos: dict, ruta_salida: str) -> str:
        """
        Genera un Excel con dos hojas: Resumen y Detalle de actividades.
        :param datos: mismo dict que generar_pdf().
        :param ruta_salida: ruta completa del archivo .xlsx de destino.
        :returns: ruta_salida.
        """
        os.makedirs(os.path.dirname(ruta_salida) if os.path.dirname(ruta_salida) else ".", exist_ok=True)

        wb = openpyxl.Workbook()

        # Hoja 1 — Resumen
        ws_res = wb.active
        ws_res.title = "Resumen"
        self._excel_hoja_resumen(ws_res, datos)

        # Hoja 2 — Detalle
        ws_det = wb.create_sheet("Detalle de Actividades")
        self._excel_hoja_detalle(ws_det, datos)

        wb.save(ruta_salida)
        return ruta_salida

    # ------------------------------------------------------------------
    # Bloques de construcción del PDF
    # ------------------------------------------------------------------

    def _encabezado(self, datos: dict, estilos: dict) -> list:
        titulo   = datos.get("titulo", "Informe Consolidado de Actividades")
        subtitulo = self._subtitulo_periodo(datos)
        asic_nombre = datos.get("asic_nombre") or datos.get("cabecera", {}).get("asic_nombre", "")
        alcance = f"ASIC: {asic_nombre}" if asic_nombre else "Todos los ASIC"

        return [
            Paragraph("CORPORACIÓN DE SALUD DEL ESTADO TÁCHIRA", estilos["institucion"]),
            Paragraph("Dirección de Programas y Proyectos de Promoción de Salud (DPPPS)", estilos["dependencia"]),
            Spacer(1, 0.2*cm),
            Paragraph(titulo, estilos["titulo"]),
            Paragraph(f"{subtitulo}  |  {alcance}", estilos["subtitulo_gris"]),
        ]

    def _tarjetas_resumen(self, datos: dict, estilos: dict) -> list:
        total_act  = datos.get("total_actividades",  0)
        total_part = datos.get("total_participantes", 0)

        tabla_data = [
            [
                Paragraph("TOTAL ACTIVIDADES", estilos["chip_label"]),
                Paragraph("TOTAL PARTICIPANTES", estilos["chip_label"]),
            ],
            [
                Paragraph(str(total_act),  estilos["chip_valor"]),
                Paragraph(str(total_part), estilos["chip_valor"]),
            ],
        ]
        tabla = Table(tabla_data, colWidths=["50%", "50%"])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), AZUL_OSCURO),
            ("BACKGROUND", (1, 0), (1, -1), ROJO_ACENTO),
            ("TEXTCOLOR",  (0, 0), (-1, -1), BLANCO),
            ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
            ("VALIGN",     (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [None]),
            ("TOPPADDING",  (0, 0), (-1, -1), 12),
            ("BOTTOMPADDING",(0, 0), (-1, -1), 12),
            ("ROUNDEDCORNERS", [4]),
        ]))
        return [tabla]

    def _tabla_resumen_asic(self, resumen: list, estilos: dict) -> list:
        return self._tabla_resumen_generica(
            titulo="Resumen por ASIC",
            resumen=resumen,
            estilos=estilos,
            color_cabecera=AZUL_OSCURO,
        )

    def _tabla_resumen_tipo(self, resumen: list, estilos: dict) -> list:
        return self._tabla_resumen_generica(
            titulo="Resumen por Tipo de Actividad",
            resumen=resumen,
            estilos=estilos,
            color_cabecera=AZUL_MEDIO,
        )

    def _tabla_resumen_generica(self, titulo: str, resumen: list,
                                 estilos: dict, color_cabecera) -> list:
        elems = [Paragraph(titulo, estilos["subtitulo"])]

        cab = [
            Paragraph("Nombre", estilos["th"]),
            Paragraph("Actividades", estilos["th"]),
            Paragraph("Participantes", estilos["th"]),
        ]
        filas = [cab]
        for i, fila in enumerate(resumen):
            nombre = fila.get("nombre") or fila.get("asic_nombre", "")
            filas.append([
                Paragraph(nombre, estilos["td"]),
                Paragraph(str(fila["total_actividades"]),  estilos["td_centro"]),
                Paragraph(str(fila["total_participantes"]), estilos["td_centro"]),
            ])

        tabla = Table(filas, colWidths=["60%", "20%", "20%"])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), color_cabecera),
            ("TEXTCOLOR",  (0, 0), (-1, 0), BLANCO),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, 0), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [BLANCO, GRIS_CLARO]),
            ("GRID",       (0, 0), (-1, -1), 0.5, GRIS_BORDE),
            ("TOPPADDING",    (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING",   (0, 0), (-1, -1), 6),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 6),
        ]))
        elems.append(tabla)
        return elems

    def _tabla_detalle(self, actividades: list, estilos: dict) -> list:
        elems = [Paragraph("Detalle de Actividades", estilos["subtitulo"])]

        cab = [
            Paragraph("#",             estilos["th"]),
            Paragraph("ASIC",          estilos["th"]),
            Paragraph("Tipo",          estilos["th"]),
            Paragraph("Fecha Inicio",  estilos["th"]),
            Paragraph("Fecha Fin",     estilos["th"]),
            Paragraph("Lugar",         estilos["th"]),
            Paragraph("Participantes", estilos["th"]),
            Paragraph("Descripción",   estilos["th"]),
        ]
        filas = [cab]
        for i, act in enumerate(actividades, start=1):
            f_ini = str(act.get("fecha_inicio", ""))
            f_fin = str(act.get("fecha_fin",   ""))
            filas.append([
                Paragraph(str(i),                         estilos["td_centro"]),
                Paragraph(act.get("asic_nombre", ""),     estilos["td_pequeño"]),
                Paragraph(act.get("tipo_nombre", ""),     estilos["td_pequeño"]),
                Paragraph(f_ini,                          estilos["td_centro"]),
                Paragraph(f_fin,                          estilos["td_centro"]),
                Paragraph(act.get("lugar", ""),           estilos["td_pequeño"]),
                Paragraph(str(act.get("participantes",0)), estilos["td_centro"]),
                Paragraph(act.get("descripcion", "")[:120], estilos["td_pequeño"]),
            ])

        col_widths = ["4%", "13%", "12%", "9%", "9%", "13%", "8%", "32%"]
        tabla = Table(filas, colWidths=col_widths, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), AZUL_OSCURO),
            ("TEXTCOLOR",  (0, 0), (-1, 0), BLANCO),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, 0), 8),
            ("FONTSIZE",   (0, 1), (-1, -1), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [BLANCO, GRIS_CLARO]),
            ("GRID",       (0, 0), (-1, -1), 0.4, GRIS_BORDE),
            ("VALIGN",     (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING",    (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING",   (0, 0), (-1, -1), 4),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 4),
        ]))
        elems.append(tabla)
        return elems

    # ------------------------------------------------------------------
    # Hojas Excel
    # ------------------------------------------------------------------

    def _excel_hoja_resumen(self, ws, datos: dict):
        """Hoja 1: encabezado + tarjetas totales + resúmenes por ASIC y tipo."""
        h_fill  = PatternFill("solid", fgColor="FF0F3460")
        h_font  = Font(color="FFFFFFFF", bold=True, size=11)
        r_fill  = PatternFill("solid", fgColor="FFF4F6F8")
        border  = Border(
            left=Side(style="thin", color="FFCCCCCC"),
            right=Side(style="thin", color="FFCCCCCC"),
            top=Side(style="thin", color="FFCCCCCC"),
            bottom=Side(style="thin", color="FFCCCCCC"),
        )

        titulo = datos.get("titulo", "Informe Consolidado")
        ws.merge_cells("A1:F1")
        ws["A1"] = titulo
        ws["A1"].font = Font(bold=True, size=14, color="FF0F3460")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        subtitulo = self._subtitulo_periodo(datos)
        ws.merge_cells("A2:F2")
        ws["A2"] = subtitulo
        ws["A2"].font = Font(size=10, color="FF888888", italic=True)
        ws["A2"].alignment = Alignment(horizontal="center")

        # Tarjetas totales
        ws["A4"] = "Total Actividades"
        ws["B4"] = datos.get("total_actividades", 0)
        ws["D4"] = "Total Participantes"
        ws["E4"] = datos.get("total_participantes", 0)
        for cell in ["A4", "D4"]:
            ws[cell].font = Font(bold=True, color="FF0F3460")
        for cell in ["B4", "E4"]:
            ws[cell].font = Font(bold=True, size=14)

        # Resumen por ASIC
        fila = 6
        ws.cell(fila, 1, "RESUMEN POR ASIC").font = Font(bold=True, size=11, color="FF0F3460")
        fila += 1
        for col, cab in enumerate(["ASIC", "Actividades", "Participantes"], 1):
            c = ws.cell(fila, col, cab)
            c.fill = h_fill; c.font = h_font
            c.alignment = Alignment(horizontal="center")
            c.border = border
        fila += 1
        for i, row in enumerate(datos.get("resumen_por_asic", [])):
            nombre = row.get("nombre") or row.get("asic_nombre", "")
            for col, val in enumerate([nombre, row["total_actividades"], row["total_participantes"]], 1):
                c = ws.cell(fila, col, val)
                c.border = border
                if i % 2 == 0:
                    c.fill = r_fill
            fila += 1

        fila += 1
        ws.cell(fila, 1, "RESUMEN POR TIPO DE ACTIVIDAD").font = Font(bold=True, size=11, color="FF0F3460")
        fila += 1
        for col, cab in enumerate(["Tipo de Actividad", "Actividades", "Participantes"], 1):
            c = ws.cell(fila, col, cab)
            c.fill = h_fill; c.font = h_font
            c.alignment = Alignment(horizontal="center")
            c.border = border
        fila += 1
        for i, row in enumerate(datos.get("resumen_por_tipo", [])):
            for col, val in enumerate([row["nombre"], row["total_actividades"], row["total_participantes"]], 1):
                c = ws.cell(fila, col, val)
                c.border = border
                if i % 2 == 0:
                    c.fill = r_fill
            fila += 1

        # Ajustar anchos
        for col, w in zip(["A", "B", "C", "D", "E", "F"], [45, 15, 15, 30, 15, 15]):
            ws.column_dimensions[col].width = w

    def _excel_hoja_detalle(self, ws, datos: dict):
        """Hoja 2: todas las actividades una por fila."""
        h_fill = PatternFill("solid", fgColor="FF0F3460")
        h_font = Font(color="FFFFFFFF", bold=True, size=10)
        r_fill = PatternFill("solid", fgColor="FFF4F6F8")
        border = Border(
            left=Side(style="thin", color="FFCCCCCC"),
            right=Side(style="thin", color="FFCCCCCC"),
            top=Side(style="thin", color="FFCCCCCC"),
            bottom=Side(style="thin", color="FFCCCCCC"),
        )

        cabeceras = ["#", "ASIC", "Tipo", "Fecha Inicio", "Fecha Fin",
                     "Lugar", "Participantes", "Descripción"]
        for col, cab in enumerate(cabeceras, 1):
            c = ws.cell(1, col, cab)
            c.fill = h_fill; c.font = h_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
        ws.row_dimensions[1].height = 22

        for i, act in enumerate(datos.get("actividades", []), start=1):
            fila_idx = i + 1
            valores = [
                i,
                act.get("asic_nombre", ""),
                act.get("tipo_nombre", ""),
                str(act.get("fecha_inicio", "")),
                str(act.get("fecha_fin", "")),
                act.get("lugar", ""),
                act.get("participantes", 0),
                act.get("descripcion", ""),
            ]
            for col, val in enumerate(valores, 1):
                c = ws.cell(fila_idx, col, val)
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=(col == 8))
                if i % 2 == 0:
                    c.fill = r_fill

        # Anchos de columna
        anchos = [5, 25, 22, 13, 13, 25, 13, 60]
        for col_idx, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = w

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _subtitulo_periodo(self, datos: dict) -> str:
        """Construye el texto del período a partir del dict de datos."""
        cabecera = datos.get("cabecera", {})
        tipo = cabecera.get("periodo_tipo") or datos.get("tipo_periodo", "")
        anio = cabecera.get("periodo_anio") or datos.get("anio_periodo", "")
        f_ini = cabecera.get("periodo_fecha_inicio") or datos.get("fecha_inicio", "")
        f_fin = cabecera.get("periodo_fecha_fin")    or datos.get("fecha_fin", "")

        if tipo and anio:
            return f"Período {tipo.capitalize()} {anio}"
        if f_ini and f_fin:
            return f"Del {f_ini} al {f_fin}"
        return ""

    @staticmethod
    def _estilos() -> dict:
        base = getSampleStyleSheet()
        return {
            "institucion": ParagraphStyle("institucion", parent=base["Normal"],
                fontSize=10, textColor=AZUL_OSCURO, fontName="Helvetica-Bold",
                alignment=TA_CENTER),
            "dependencia": ParagraphStyle("dependencia", parent=base["Normal"],
                fontSize=8, textColor=AZUL_MEDIO, alignment=TA_CENTER),
            "titulo": ParagraphStyle("titulo", parent=base["Title"],
                fontSize=14, textColor=AZUL_OSCURO, fontName="Helvetica-Bold",
                alignment=TA_CENTER, spaceAfter=2),
            "subtitulo_gris": ParagraphStyle("subtitulo_gris", parent=base["Normal"],
                fontSize=9, textColor=colors.HexColor("#888888"), alignment=TA_CENTER),
            "subtitulo": ParagraphStyle("subtitulo", parent=base["Heading2"],
                fontSize=10, textColor=AZUL_OSCURO, fontName="Helvetica-Bold",
                spaceBefore=6, spaceAfter=4),
            "normal": ParagraphStyle("normal", parent=base["Normal"],
                fontSize=9, leading=13),
            "th": ParagraphStyle("th", parent=base["Normal"],
                fontSize=8, textColor=BLANCO, fontName="Helvetica-Bold",
                alignment=TA_CENTER),
            "td": ParagraphStyle("td", parent=base["Normal"],
                fontSize=8, leading=11),
            "td_centro": ParagraphStyle("td_centro", parent=base["Normal"],
                fontSize=8, alignment=TA_CENTER, leading=11),
            "td_pequeño": ParagraphStyle("td_pequeño", parent=base["Normal"],
                fontSize=7, leading=10),
            "chip_label": ParagraphStyle("chip_label", parent=base["Normal"],
                fontSize=8, textColor=BLANCO, fontName="Helvetica-Bold",
                alignment=TA_CENTER),
            "chip_valor": ParagraphStyle("chip_valor", parent=base["Normal"],
                fontSize=22, textColor=BLANCO, fontName="Helvetica-Bold",
                alignment=TA_CENTER),
            "pie": ParagraphStyle("pie", parent=base["Normal"],
                fontSize=7, textColor=colors.HexColor("#aaaaaa"),
                alignment=TA_RIGHT),
        }
